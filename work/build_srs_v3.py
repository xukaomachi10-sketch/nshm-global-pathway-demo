from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "SRS_HeThong_QuanLy_HoSo_Du_Hoc_NSHM_v3_Pro.docx"
UI_DIR = ROOT / "work" / "ui_crops"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

COLORS = {
    "navy": "0B2545",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "red": "C8102E",
    "ink": "20262E",
    "muted": "667085",
    "line": "CBD5E1",
    "table_header": "E8EEF5",
    "table_alt": "F8FAFC",
    "callout": "F4F6F9",
    "success": "EAF6EE",
    "warning": "FFF7E6",
    "risk": "FDECEC",
    "white": "FFFFFF",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}, got {sum(widths_dxa)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(index, len(widths_dxa) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color: str = "B8C4D1", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_run_font(run, name="Calibri", size=None, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_paragraph_border(paragraph, side: str, color: str, size: int = 8, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color="20262E", size=8.8, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = "Table Text"
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text)) if not paragraph.text else paragraph.runs[0]
    if paragraph.text and paragraph.runs:
        paragraph.runs[0].text = str(text)
        run = paragraph.runs[0]
    set_run_font(run, size=size, bold=bold, color=color)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
    *,
    header_fill: str = COLORS["table_header"],
    alt_rows: bool = False,
    font_size: float = 8.8,
    center_columns: Iterable[int] = (),
) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    keep_row_together(header)
    for index, text in enumerate(headers):
        set_cell_shading(header.cells[index], header_fill)
        set_cell_text(
            header.cells[index],
            text,
            bold=True,
            color=COLORS["navy"],
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
        )
    center_columns = set(center_columns)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        if alt_rows and row_index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, COLORS["table_alt"])
        for col_index, value in enumerate(values):
            set_cell_text(
                row.cells[col_index],
                value,
                size=font_size,
                align=WD_ALIGN_PARAGRAPH.CENTER if col_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
            )
    doc.add_paragraph(style="Table Spacing")
    return table


def add_field(paragraph, field_code: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=COLORS["muted"])


def add_callout(doc: Document, label: str, text: str, kind: str = "info") -> None:
    fills = {"info": COLORS["callout"], "success": COLORS["success"], "warning": COLORS["warning"], "risk": COLORS["risk"]}
    accents = {"info": COLORS["blue"], "success": "2F855A", "warning": "B7791F", "risk": COLORS["red"]}
    paragraph = doc.add_paragraph(style="Callout")
    set_paragraph_shading(paragraph, fills[kind])
    set_paragraph_border(paragraph, "left", accents[kind], size=18, space=8)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=accents[kind])
    text_run = paragraph.add_run(text)
    set_run_font(text_run, color=COLORS["ink"])


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True, color=COLORS["navy"])
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, color=COLORS["ink"])
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=COLORS["ink"])
    return paragraph


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="SRS Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = level
    num_id = num_pr.get_or_add_numId()
    num_id.val = 10
    run = paragraph.add_run(text)
    set_run_font(run, color=COLORS["ink"])


def add_number(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="SRS Number")
    paragraph.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = level
    num_id = num_pr.get_or_add_numId()
    num_id.val = 11
    run = paragraph.add_run(text)
    set_run_font(run, color=COLORS["ink"])


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=9, italic=True, color=COLORS["muted"])


def add_screen(doc: Document, filename: str, caption: str) -> None:
    path = UI_DIR / filename
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(6.3))
    add_caption(doc, caption)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_requirement_module(
    doc: Document,
    section_no: str,
    title: str,
    objective: str,
    requirements: Sequence[tuple[str, str, str, str]],
    rules: Sequence[str],
    screenshot: tuple[str, str] | None = None,
) -> None:
    add_heading(doc, f"{section_no}. {title}", 2)
    objective_paragraph = add_body(doc, objective)
    objective_paragraph.paragraph_format.keep_with_next = True
    if screenshot:
        add_screen(doc, screenshot[0], screenshot[1])
    add_table(
        doc,
        ["ID", "Yêu cầu", "Ưu tiên", "Tiêu chí chấp nhận / ghi chú"],
        requirements,
        [900, 4950, 850, 2660],
        alt_rows=True,
        center_columns=(0, 2),
        font_size=8.45,
    )
    add_heading(doc, "Quy tắc nghiệp vụ chính", 3)
    for rule in rules:
        add_bullet(doc, rule)


def add_custom_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part.element
    for abstract_id, num_id, fmt, text in ((10, 10, "bullet", "•"), (11, 11, "decimal", "%1.")):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), fmt)
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), text if fmt == "bullet" else f"%{level + 1}.")
            lvl_jc = OxmlElement("w:lvlJc")
            lvl_jc.set(qn("w:val"), "left")
            p_pr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(540 + level * 360))
            tabs.append(tab)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(540 + level * 360))
            ind.set(qn("w:hanging"), "270")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "80")
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
            p_pr.extend([tabs, ind, spacing])
            lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
            if fmt == "bullet":
                r_pr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Calibri")
                fonts.set(qn("w:hAnsi"), "Calibri")
                r_pr.append(fonts)
                lvl.append(r_pr)
            abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(COLORS["ink"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, COLORS["blue"], 18, 10),
        2: (13, COLORS["blue"], 14, 7),
        3: (12, COLORS["dark_blue"], 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, base, size, after, line in (
        ("Table Text", "Normal", 8.8, 0, 1.05),
        ("Table Spacing", "Normal", 2, 4, 1.0),
        ("Callout", "Normal", 10.5, 8, 1.15),
        ("SRS Bullet", "Normal", 11, 4, 1.25),
        ("SRS Number", "Normal", 11, 4, 1.25),
    ):
        if name not in styles:
            styles.add_style(name, 1)
        style = styles[name]
        style.base_style = styles[base]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(COLORS["muted"])
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("NSHM | Hệ thống Quản lý Hồ sơ Du học")
    set_run_font(left, size=8.5, bold=True, color=COLORS["muted"])
    right = p.add_run("\tSRS v3.0")
    set_run_font(right, size=8.5, color=COLORS["muted"])

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fr = fp.add_run("Nội bộ - Phòng Hợp tác Quốc tế | Trang ")
    set_run_font(fr, size=9, color=COLORS["muted"])
    add_field(fp, "PAGE", "1")


def build_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("PHÒNG HỢP TÁC QUỐC TẾ & TƯ VẤN HƯỚNG NGHIỆP")
    set_run_font(r, size=10.5, bold=True, color=COLORS["red"])

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(34)
    title.paragraph_format.space_after = Pt(8)
    tr = title.add_run("SOFTWARE REQUIREMENTS\nSPECIFICATION")
    set_run_font(tr, size=28, bold=True, color=COLORS["navy"])

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("HỆ THỐNG QUẢN LÝ HỒ SƠ DU HỌC")
    set_run_font(sr, size=18, bold=True, color=COLORS["blue"])

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(2)
    lead.paragraph_format.space_after = Pt(28)
    lr = lead.add_run("Bản đặc tả nghiệp vụ và phần mềm - Baseline bàn giao cho đội IT")
    set_run_font(lr, size=12, color=COLORS["muted"])

    set_paragraph_border(subtitle, "bottom", COLORS["red"], size=18, space=10)

    metadata = [
        ("Đơn vị sở hữu", "Phòng Hợp tác Quốc tế - Trường Ngôi Sao Hoàng Mai"),
        ("Phiên bản", "v3.0 - Professional SRS"),
        ("Trạng thái", "Đề xuất làm baseline để IT phân tích, ước lượng và phát triển"),
        ("Quy mô mục tiêu", "1.000 hồ sơ/năm học; kiến trúc sẵn sàng mở rộng 5.000 hồ sơ tích lũy"),
        ("Ngày cập nhật", "30/06/2026"),
        ("Mức độ bảo mật", "Nội bộ - Có dữ liệu cá nhân và hồ sơ học sinh"),
    ]
    add_table(doc, ["Thông tin", "Nội dung"], metadata, [2100, 7260], font_size=9.5)
    add_callout(
        doc,
        "Mục tiêu thiết kế",
        "Chuyên viên nhìn được toàn cảnh, biết ngay việc cần làm tiếp theo, xử lý tài liệu không thất lạc và kiểm soát deadline của 1.000 hồ sơ một cách nhất quán.",
        "success",
    )
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    nr = note.add_run("Nguồn đầu vào: SRS v2.0 và bộ giao diện UI đã cập nhật.")
    set_run_font(nr, size=9.5, italic=True, color=COLORS["muted"])
    add_page_break(doc)


def build_front_matter(doc: Document) -> None:
    add_heading(doc, "KIỂM SOÁT TÀI LIỆU", 1)
    add_callout(
        doc,
        "Cách dùng tài liệu",
        "Các yêu cầu có mã FR/NFR là đầu vào truy vết cho thiết kế, backlog và kiểm thử. Từ Must là phạm vi bắt buộc của baseline; thay đổi Must phải được chủ nghiệp vụ phê duyệt.",
        "info",
    )
    add_heading(doc, "Lịch sử phiên bản", 2)
    add_table(
        doc,
        ["Phiên bản", "Ngày", "Mô tả thay đổi", "Chủ trì"],
        [
            ("v2.0", "30/06/2026", "Bổ sung bộ giao diện UI, module website công khai và CMS.", "Phòng HTQT"),
            ("v3.0", "30/06/2026", "Chuẩn hóa SRS: yêu cầu có ID, RBAC, workflow, data model, NFR đo được, UAT, roadmap và quyết định mở.", "Phòng HTQT / Codex"),
        ],
        [1100, 1200, 5350, 1710],
        center_columns=(0, 1),
    )
    add_heading(doc, "Phê duyệt baseline", 2)
    add_table(
        doc,
        ["Vai trò phê duyệt", "Họ tên", "Trạng thái", "Ngày / chữ ký"],
        [
            ("Trưởng phòng HTQT", "", "Chờ phê duyệt", ""),
            ("Đại diện IT / Nhà cung cấp", "", "Chờ xác nhận khả thi", ""),
            ("Đại diện BGH", "", "Chờ phê duyệt phạm vi", ""),
            ("Đại diện bảo mật / dữ liệu", "", "Chờ rà soát", ""),
        ],
        [2500, 2300, 2200, 2360],
        center_columns=(2,),
    )
    add_heading(doc, "Quy ước ưu tiên", 2)
    add_table(
        doc,
        ["Mức", "Ý nghĩa", "Quy tắc"],
        [
            ("Must", "Bắt buộc", "Thiếu yêu cầu này thì hệ thống không thể nghiệm thu baseline."),
            ("Should", "Nên có", "Giá trị cao; có thể lùi sang đợt phát hành kế tiếp nếu được phê duyệt."),
            ("Could", "Có thể có", "Tối ưu trải nghiệm; không chặn go-live."),
            ("Out", "Ngoài phạm vi", "Không triển khai trong baseline; chỉ giữ điểm mở rộng."),
        ],
        [1300, 1700, 6360],
        center_columns=(0,),
    )
    add_heading(doc, "Giả định và quyết định cần chốt", 2)
    add_table(
        doc,
        ["ID", "Nội dung", "Giả định tạm thời", "Chủ trì chốt"],
        [
            ("OD-01", "Nền tảng lưu trữ file", "Thiết kế lớp tích hợp hỗ trợ Google Drive hoặc Microsoft 365; chọn một nhà cung cấp cho go-live.", "IT + HTQT"),
            ("OD-02", "Nguồn dữ liệu học sinh", "Mã học sinh từ SIS là khóa nghiệp vụ chính; cho phép nhập tạm nếu SIS chưa sẵn sàng.", "IT + Giáo vụ"),
            ("OD-03", "Kênh đăng nhập", "Nhân sự dùng SSO của trường; học sinh/CMHS dùng tài khoản trường hoặc OTP theo kiến trúc được duyệt.", "IT"),
            ("OD-04", "Chính sách lưu giữ", "Thời hạn lưu hồ sơ, log và file là cấu hình; không xóa vật lý trước khi chính sách được phê duyệt.", "BGH + IT"),
            ("OD-05", "Website và CMS", "Là Phase 3; không chặn go-live Core CRM nếu nguồn lực hạn chế.", "BGH + HTQT"),
            ("OD-06", "Email/Teams/SMS", "Email và thông báo trong hệ thống là baseline; Teams/SMS là tùy chọn tích hợp.", "IT + HTQT"),
        ],
        [850, 2550, 4300, 1660],
        center_columns=(0,),
        font_size=8.6,
    )

    add_page_break(doc)
    add_heading(doc, "MỤC LỤC", 1)
    toc_rows = [
        ("1", "Tổng quan và mục tiêu"),
        ("2", "Phạm vi, tác nhân và phân quyền"),
        ("3", "Quy trình nghiệp vụ và trạng thái"),
        ("4", "Yêu cầu chức năng"),
        ("5", "Mô hình dữ liệu và quản trị tài liệu"),
        ("6", "Yêu cầu giao diện và trải nghiệm"),
        ("7", "Tích hợp và API"),
        ("8", "Yêu cầu phi chức năng"),
        ("9", "Bảo mật, riêng tư và audit"),
        ("10", "Báo cáo và KPI"),
        ("11", "Kiểm thử, UAT và tiêu chí nghiệm thu"),
        ("12", "Triển khai, di chuyển dữ liệu và vận hành"),
        ("13", "Roadmap và phạm vi phát hành"),
        ("Phụ lục A", "Từ điển trạng thái"),
        ("Phụ lục B", "Ma trận truy vết"),
        ("Phụ lục C", "Checklist hồ sơ mẫu"),
    ]
    add_table(doc, ["Mục", "Nội dung"], toc_rows, [1400, 7960], font_size=10)
    add_body(doc, "Lưu ý: số trang có thể thay đổi khi đội IT bổ sung kiến trúc, sơ đồ hoặc biên bản phê duyệt. Mã yêu cầu không thay đổi theo số trang.")
    add_page_break(doc)


def build_overview(doc: Document) -> None:
    add_heading(doc, "1. TỔNG QUAN VÀ MỤC TIÊU", 1)
    add_heading(doc, "1.1. Mục đích", 2)
    add_body(doc, "Tài liệu xác định yêu cầu nghiệp vụ và phần mềm cho hệ thống quản lý toàn bộ hành trình du học của học sinh, từ đăng ký tư vấn đến kết quả ứng tuyển và nhập học. Đối tượng đọc gồm chủ nghiệp vụ, đội phát triển, kiểm thử, vận hành, bảo mật và đơn vị nghiệm thu.")
    add_heading(doc, "1.2. Bối cảnh nghiệp vụ", 2)
    add_body(doc, "Phòng Hợp tác Quốc tế cần thay thế cách theo dõi phân tán bằng bảng tính, tin nhắn và thư mục rời rạc. Hệ thống phải biến mỗi hồ sơ thành một case có chủ sở hữu, trạng thái, việc tiếp theo, deadline, checklist và lịch sử thay đổi rõ ràng.")
    for item in (
        "Học sinh đăng ký để được tư vấn; phòng HTQT xác minh học sinh có thực sự làm hồ sơ hay chỉ tham khảo.",
        "Dữ liệu học thuật gồm điểm học kỳ, điểm môn, IELTS, SAT và các bài thi chuẩn hóa khác.",
        "Thành tích, chứng chỉ, CLB và hoạt động ngoại khóa được nộp kèm minh chứng và tự động đưa vào đúng thư mục.",
        "Chuyên viên nhìn toàn cảnh 1.000 hồ sơ, ưu tiên đúng học sinh và không bỏ lỡ deadline.",
        "Essay, thư giới thiệu và tài liệu ứng tuyển được quản lý theo phiên bản, người phụ trách và trạng thái duyệt.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "1.3. Mục tiêu đo được", 2)
    add_table(
        doc,
        ["Mục tiêu", "Chỉ số", "Ngưỡng nghiệm thu sau ổn định"],
        [
            ("Một hồ sơ duy nhất", "Tỷ lệ trùng StudentID", "0 hồ sơ trùng do cùng Mã HS; mọi lượt đăng ký lặp được liên kết."),
            ("Không bỏ lỡ deadline", "Deadline bắt buộc quá hạn không cảnh báo", "0 trường hợp do hệ thống không tạo hoặc không hiển thị cảnh báo."),
            ("Hồ sơ có chủ sở hữu", "Case đang hoạt động chưa phân công", "Dưới 1% và không quá 1 ngày làm việc."),
            ("Tài liệu có kiểm soát", "File không xác định học sinh / hạng mục", "Dưới 0,5%; có hàng đợi xử lý ngoại lệ."),
            ("Vận hành thuận tiện", "Thời gian tìm hồ sơ", "Dưới 2 giây ở p95 trong điều kiện tải chuẩn."),
            ("Minh bạch quản trị", "Thao tác nhạy cảm có audit", "100% sự kiện bắt buộc được ghi log."),
        ],
        [2300, 2700, 4360],
        font_size=8.8,
    )
    add_heading(doc, "1.4. Phạm vi", 2)
    add_table(
        doc,
        ["Trong phạm vi", "Ngoài phạm vi baseline"],
        [
            ("Đăng ký, chống trùng, xác minh nhu cầu, phân công và đặt lịch tư vấn.", "Nộp hồ sơ trực tiếp vào cổng của các trường đại học."),
            ("Student 360, điểm học kỳ, điểm chuẩn hóa, định hướng du học.", "Thanh toán học phí, quản lý tài chính kế toán hoặc hợp đồng dịch vụ."),
            ("Kho minh chứng, tài liệu, essay, thư giới thiệu, checklist và phiên bản.", "AI tự viết essay, tự chấm cơ hội đỗ hoặc ra quyết định thay chuyên viên."),
            ("Pipeline trường/ngành, deadline, phỏng vấn, kết quả, học bổng.", "Thay thế SIS/LMS/ERP của nhà trường."),
            ("Dashboard, báo cáo, thông báo, audit, import/export và tích hợp.", "Lưu mật khẩu cổng tuyển sinh của trường đại học nếu chưa có kiểm soát bảo mật riêng."),
            ("Portal học sinh/CMHS; website và CMS theo roadmap.", "Ứng dụng mobile native trong baseline."),
        ],
        [4680, 4680],
        alt_rows=True,
        font_size=8.8,
    )
    add_heading(doc, "1.5. Thuật ngữ", 2)
    add_table(
        doc,
        ["Thuật ngữ", "Định nghĩa"],
        [
            ("Case / Hồ sơ", "Bản ghi quản lý hành trình du học của một học sinh trong một chu kỳ tư vấn."),
            ("Student 360", "Màn hình tổng hợp hồ sơ học sinh, điểm, tài liệu, ứng tuyển, lịch sử và cảnh báo."),
            ("Evidence", "Minh chứng cho chứng chỉ, giải thưởng, CLB, sự kiện hoặc hoạt động ngoại khóa."),
            ("Document item", "Một hạng mục tài liệu phải chuẩn bị, có người phụ trách, trạng thái, deadline và file phiên bản."),
            ("Application", "Một nguyện vọng ứng tuyển duy nhất theo trường + chương trình + kỳ nhập học."),
            ("Checklist", "Danh sách hạng mục bắt buộc hoặc tùy chọn để hoàn thành hồ sơ."),
            ("RBAC", "Phân quyền theo vai trò; kết hợp phạm vi bản ghi được phân công."),
            ("RPO / RTO", "Mức mất dữ liệu tối đa chấp nhận / thời gian phục hồi mục tiêu."),
        ],
        [1900, 7460],
    )


def build_scope_roles(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "2. PHẠM VI, TÁC NHÂN VÀ PHÂN QUYỀN", 1)
    add_heading(doc, "2.1. Tác nhân", 2)
    add_table(
        doc,
        ["Tác nhân", "Mục tiêu sử dụng", "Phạm vi dữ liệu mặc định"],
        [
            ("Học sinh", "Đăng ký, cập nhật hồ sơ, nộp file, xem việc cần làm và lịch hẹn.", "Chỉ hồ sơ của chính mình."),
            ("CMHS", "Theo dõi tiến độ, lịch hẹn và bổ sung tài liệu khi được ủy quyền.", "Hồ sơ con đã liên kết và được đồng ý."),
            ("Chuyên viên", "Quản lý case được giao, tư vấn, review tài liệu, theo dõi ứng tuyển.", "Case đang/đã được phân công."),
            ("Trưởng phòng HTQT", "Điều phối nguồn lực, xem toàn bộ case, phê duyệt ngoại lệ và báo cáo.", "Toàn bộ đơn vị HTQT."),
            ("BGH", "Xem dashboard và báo cáo tổng hợp; xem chi tiết khi có quyền rõ ràng.", "Tổng hợp; chi tiết theo phê duyệt."),
            ("Giáo viên / Người giới thiệu", "Cung cấp thư giới thiệu hoặc xác nhận hoạt động.", "Chỉ hạng mục được mời; nội dung có thể confidential."),
            ("Admin hệ thống", "Quản trị tài khoản, cấu hình, tích hợp, log và hỗ trợ kỹ thuật.", "Không mặc định sửa nội dung nghiệp vụ."),
            ("Admin CMS", "Soạn, duyệt và xuất bản nội dung website.", "Module CMS và media."),
            ("Hệ thống tích hợp", "Đồng bộ SIS, SSO, lịch, email, storage và BI.", "Tài khoản dịch vụ theo least privilege."),
        ],
        [2000, 4100, 3260],
        alt_rows=True,
        font_size=8.6,
    )
    add_heading(doc, "2.2. Nguyên tắc phân quyền", 2)
    for item in (
        "Quyền hiệu lực = quyền theo vai trò + phạm vi bản ghi + trạng thái hồ sơ + cờ confidential.",
        "Admin kỹ thuật không mặc định được đọc essay, thư giới thiệu và tài liệu tài chính; truy cập hỗ trợ phải có lý do và audit.",
        "Mọi export chứa dữ liệu cá nhân cần quyền riêng; file export có thời hạn tải và được ghi log.",
        "Thay đổi phân công không làm mất lịch sử quyền sở hữu; quyền người cũ kết thúc theo cấu hình bàn giao.",
        "BGH ưu tiên xem dữ liệu tổng hợp; quyền drill-down đến hồ sơ cá nhân cần được phê duyệt.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "2.3. Ma trận quyền cấp cao", 2)
    add_table(
        doc,
        ["Module", "Học sinh / CMHS", "Chuyên viên", "Trưởng phòng", "BGH", "Admin"],
        [
            ("Đăng ký & lịch tư vấn", "Tạo / xem của mình", "Xử lý được giao", "Toàn quyền nghiệp vụ", "Xem tổng hợp", "Cấu hình"),
            ("Student 360", "Xem bản rút gọn", "Xem/sửa case được giao", "Xem/sửa toàn đơn vị", "Xem theo quyền", "Hỗ trợ có audit"),
            ("Điểm & chứng chỉ", "Nộp / đề xuất sửa", "Xác minh / sửa", "Phê duyệt ngoại lệ", "Xem tổng hợp", "Cấu hình danh mục"),
            ("Evidence & file", "Upload của mình", "Review case được giao", "Toàn đơn vị", "Không mặc định", "Quản trị storage"),
            ("Ứng tuyển", "Xem trạng thái", "Tạo/cập nhật", "Toàn quyền", "Xem tổng hợp", "Cấu hình master"),
            ("Essay / thư giới thiệu", "Theo chính sách item", "Review theo case", "Xem theo confidential", "Không mặc định", "Không mặc định"),
            ("Dashboard / báo cáo", "Cá nhân", "Phạm vi được giao", "Toàn đơn vị", "Tổng hợp", "Vận hành hệ thống"),
            ("CMS", "Xem public", "Không mặc định", "Duyệt nếu được cấp", "Xem", "Admin CMS riêng"),
        ],
        [1750, 1700, 1850, 1750, 1100, 1210],
        font_size=7.9,
        center_columns=(1, 2, 3, 4, 5),
    )
    add_callout(doc, "Điểm kiểm soát", "Mọi quyền 'xem toàn bộ hồ sơ' phải là quyền được cấp minh thị, có chủ sở hữu phê duyệt và được rà soát định kỳ.", "warning")


def build_workflows(doc: Document) -> None:
    add_heading(doc, "3. QUY TRÌNH NGHIỆP VỤ VÀ TRẠNG THÁI", 1)
    add_heading(doc, "3.1. Hành trình end-to-end", 2)
    steps = [
        ("1", "Tiếp nhận", "Nhận form / import / nhập tay; đối chiếu Mã HS và phát hiện trùng."),
        ("2", "Phân công & đặt lịch", "Giao chuyên viên; kiểm tra lịch; gửi xác nhận và nhắc hẹn."),
        ("3", "Tư vấn & xác nhận", "Ghi biên bản; kết luận Làm hồ sơ / Chưa quyết định / Không làm."),
        ("4", "Xây dựng hồ sơ", "Cập nhật mục tiêu, điểm, chứng chỉ, evidence, essay và checklist."),
        ("5", "Ứng tuyển", "Theo dõi từng trường/ngành, deadline, nộp hồ sơ và phỏng vấn."),
        ("6", "Kết quả & nhập học", "Ghi offer/reject/waitlist, học bổng, quyết định và trường nhập học."),
        ("7", "Đóng & lưu trữ", "Chốt outcome, khóa thay đổi phù hợp, lưu trữ và báo cáo."),
    ]
    add_table(doc, ["Bước", "Giai đoạn", "Kết quả bắt buộc"], steps, [800, 2000, 6560], center_columns=(0,), font_size=9)
    add_heading(doc, "3.2. Trạng thái case", 2)
    add_table(
        doc,
        ["Trạng thái", "Điều kiện vào", "Điều kiện ra / trạng thái kế"],
        [
            ("New Registration", "Có lượt đăng ký hợp lệ.", "Verified hoặc Needs Verification."),
            ("Needs Verification", "Thiếu/không khớp Mã HS hoặc nghi trùng.", "Verified sau khi liên kết/tạo Student đúng."),
            ("Counseling Scheduled", "Đã phân công và có lịch hợp lệ.", "Counseled / No-show / Rescheduled."),
            ("Decision Pending", "Đã tư vấn nhưng chưa kết luận.", "Active Case hoặc Not Proceeding."),
            ("Active Case", "Học sinh xác nhận làm hồ sơ; có owner và intake.", "Application In Progress / On Hold / Closed."),
            ("Application In Progress", "Có ít nhất một application đang chuẩn bị/đã nộp.", "Awaiting Results / Closed."),
            ("Awaiting Results", "Đã nộp và chờ kết quả.", "Offer Received / Closed."),
            ("Offer Received", "Có ít nhất một offer.", "Enrolled / Closed."),
            ("On Hold", "Tạm dừng có lý do và ngày review.", "Trở lại Active hoặc Closed."),
            ("Not Proceeding", "Học sinh xác nhận không làm hồ sơ.", "Reopened nếu có đăng ký mới được phê duyệt."),
            ("Enrolled / Closed", "Đã chốt trường nhập học hoặc kết thúc case.", "Chỉ reopen bởi người có quyền."),
        ],
        [2100, 3100, 4160],
        alt_rows=True,
        font_size=8.6,
    )
    add_heading(doc, "3.3. Chống trùng và hợp nhất", 2)
    for item in (
        "Exact match theo Mã HS là mức tin cậy cao nhất và không được tạo Student mới.",
        "Nếu thiếu Mã HS, hệ thống gợi ý nghi trùng theo họ tên chuẩn hóa + ngày sinh + lớp + điện thoại/email; người có quyền quyết định liên kết.",
        "Hợp nhất phải hiển thị bản ghi giữ lại, bản ghi nguồn, trường xung đột và kết quả dự kiến trước khi xác nhận.",
        "Sau hợp nhất, ID cũ được lưu alias để link và audit không bị đứt; file không bị sao chép lặp.",
        "Không tự động hợp nhất chỉ dựa trên tên giống nhau.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "3.4. Quy tắc ưu tiên và rủi ro", 2)
    add_table(
        doc,
        ["Mức", "Ví dụ điều kiện", "Hành vi hệ thống"],
        [
            ("Critical", "Deadline đã quá hạn; lỗi nộp; thiếu tài liệu bắt buộc sát hạn.", "Đứng đầu worklist; cảnh báo đỏ; escalates đến Trưởng phòng."),
            ("High", "Deadline <= 7 ngày; học sinh chưa phản hồi; review bị trả lại nhiều lần.", "Cảnh báo nổi bật; tạo task và nhắc hằng ngày."),
            ("Medium", "Deadline 8-30 ngày; checklist chưa đạt ngưỡng.", "Hiển thị cam; đưa vào kế hoạch tuần."),
            ("Low", "Không có deadline gần và checklist đúng tiến độ.", "Hiển thị xanh; theo dõi định kỳ."),
        ],
        [1300, 4300, 3760],
        center_columns=(0,),
    )
    add_heading(doc, "3.5. Cấu trúc thư mục tự động", 2)
    add_table(
        doc,
        ["Thứ tự", "Thư mục", "Nội dung"],
        [
            ("00", "Profile", "Thông tin nền và hồ sơ nhận diện."),
            ("01", "Academic", "Học bạ, bảng điểm, transcript."),
            ("02", "TestScores", "IELTS, SAT, TOEFL, AP, HSK và chứng chỉ chuẩn hóa."),
            ("03", "Awards_Activities", "Giải thưởng, CLB, sự kiện, tình nguyện, trại hè."),
            ("04", "Essays", "Personal essay, supplemental essay và phiên bản."),
            ("05", "RecommendationLetters", "Thư giới thiệu; có thể giới hạn quyền xem."),
            ("06", "Applications", "Đơn, biên nhận nộp, trao đổi theo trường."),
            ("07", "Financial", "Tài liệu tài chính và học bổng."),
            ("08", "Results", "Offer, reject, waitlist, visa và quyết định nhập học."),
        ],
        [900, 2800, 5660],
        center_columns=(0,),
        alt_rows=True,
    )
    add_callout(doc, "Tính lặp an toàn", "Tác vụ tạo folder phải idempotent: chạy lại không tạo folder trùng, không đổi link hiện có và phải ghi log đồng bộ.", "info")


def build_functional_requirements(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "4. YÊU CẦU CHỨC NĂNG", 1)
    add_body(doc, "Mỗi yêu cầu có ID ổn định để ánh xạ sang user story, thiết kế, test case và biên bản nghiệm thu. Khi triển khai, IT có thể tách một FR thành nhiều ticket nhưng không được làm mất tiêu chí chấp nhận.")

    add_requirement_module(
        doc, "4.1", "Đăng nhập, tài khoản và phân quyền", "Cung cấp truy cập an toàn, ít thao tác và đúng phạm vi dữ liệu cho từng nhóm người dùng.",
        [
            ("FR-AUT-01", "Nhân sự đăng nhập bằng SSO của trường; ưu tiên Entra ID hoặc nhà cung cấp danh tính được IT phê duyệt.", "Must", "Đăng nhập thành công; tài khoản bị khóa ở IdP không vào được hệ thống."),
            ("FR-AUT-02", "Hệ thống ánh xạ nhóm danh tính sang vai trò và đơn vị; thay đổi quyền có hiệu lực theo SLA cấu hình.", "Must", "Không cần sửa từng hồ sơ khi đổi nhóm quyền."),
            ("FR-AUT-03", "Áp dụng RBAC kết hợp record scope và cờ confidential cho mọi API và màn hình.", "Must", "Truy cập URL trực tiếp vẫn bị chặn nếu không có quyền."),
            ("FR-AUT-04", "Tự động kết thúc phiên sau thời gian không hoạt động; hỗ trợ đăng xuất mọi thiết bị.", "Must", "Phiên hết hạn không gọi được API; người dùng được đưa về trang đăng nhập."),
            ("FR-AUT-05", "Tài khoản học sinh/CMHS được liên kết rõ với StudentID và có cơ chế xác minh.", "Must", "Không thể xem hồ sơ khác bằng đổi tham số URL."),
            ("FR-AUT-06", "Tài khoản dịch vụ tích hợp dùng thông tin xác thực riêng, quyền tối thiểu và có ngày rà soát.", "Must", "Có danh sách service accounts và log sử dụng."),
            ("FR-AUT-07", "Admin có thể vô hiệu hóa tài khoản, thu hồi phiên và xem lịch sử quyền.", "Must", "Quyền bị thu hồi không còn hiệu lực; sự kiện có audit."),
            ("FR-AUT-08", "Hỗ trợ ủy quyền tạm thời khi chuyên viên nghỉ, có thời hạn và người phê duyệt.", "Should", "Quyền tự hết hạn và case có lịch sử bàn giao."),
        ],
        [
            "Không dùng quyền hiển thị frontend để thay thế kiểm soát quyền tại backend.",
            "Tài khoản không hoạt động phải được rà soát và vô hiệu hóa theo chính sách IT.",
            "Chế độ hỗ trợ/impersonation nếu có phải hiển thị banner rõ ràng và ghi audit đầy đủ.",
        ],
    )

    add_requirement_module(
        doc, "4.2", "Đăng ký, chống trùng và tiếp nhận tư vấn", "Tiếp nhận nhanh nhưng không tạo nhiều hồ sơ học sinh cho cùng một người.",
        [
            ("FR-REG-01", "Cung cấp form đăng ký từ website/portal với Mã HS, thông tin liên hệ, quốc gia, ngành quan tâm, thời gian mong muốn và đồng ý xử lý dữ liệu.", "Must", "Thiếu trường bắt buộc không gửi được; thông báo thành công có mã đăng ký."),
            ("FR-REG-02", "Kiểm tra định dạng Mã HS và đối chiếu SIS/danh sách master ngay khi tiếp nhận.", "Must", "Mã không hợp lệ chuyển Needs Verification, không tự tạo Student chính thức."),
            ("FR-REG-03", "Phát hiện trùng exact theo Mã HS và gợi ý nghi trùng theo dữ liệu định danh phụ.", "Must", "Exact match liên kết hồ sơ cũ; fuzzy match cần người xác nhận."),
            ("FR-REG-04", "Một Student có nhiều CounselingRegistration và lưu đầy đủ lịch sử nguồn, thời gian, trạng thái.", "Must", "Đăng ký lặp xuất hiện trong timeline nhưng chỉ có một StudentID."),
            ("FR-REG-05", "Hiển thị hàng đợi đăng ký với bộ lọc, cảnh báo trùng, tuổi hàng đợi và thứ tự ưu tiên.", "Must", "Lọc theo ít nhất 8 tiêu chí; kết quả và tổng số nhất quán."),
            ("FR-REG-06", "Phân công một hoặc nhiều đăng ký cho chuyên viên; hỗ trợ gợi ý theo tải, quốc gia và chuyên môn.", "Must", "Bulk assign thành công; không vượt quyền; có audit."),
            ("FR-REG-07", "Cho phép ghi kết luận nhu cầu: Làm hồ sơ, Chưa quyết định, Không làm, Cần follow-up.", "Must", "Mỗi kết luận có ngày, người ghi, ghi chú và next action."),
            ("FR-REG-08", "Cho phép hợp nhất/liên kết bản ghi nghi trùng qua màn hình preview xung đột.", "Must", "Không mất lịch, file, task hoặc audit; ID cũ trở thành alias."),
            ("FR-REG-09", "Hỗ trợ import Excel có template, preview, kiểm tra lỗi theo dòng và báo cáo kết quả.", "Should", "Dòng lỗi không làm rollback dòng hợp lệ; có file kết quả import."),
            ("FR-REG-10", "Xuất danh sách theo bộ lọc hiện hành và quyền dữ liệu.", "Should", "Export không chứa cột người dùng không có quyền xem."),
        ],
        [
            "Ưu tiên hàng đợi: trùng cần xác minh > chưa phân công > chưa đặt lịch > đăng ký mới.",
            "Không xóa vật lý lượt đăng ký; hủy/duplicate là trạng thái có lý do.",
            "Thông tin liên hệ phải được chuẩn hóa để hỗ trợ tra cứu nhưng vẫn giữ giá trị gốc trong audit/import log.",
        ],
        ("UI-02_registration.png", "Hình 1 - UI-02: Danh sách đăng ký chờ tư vấn và cảnh báo trùng (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.3", "Lịch và biên bản tư vấn", "Quản lý lịch hẹn, tránh trùng, ghi nhận nội dung tư vấn và tạo việc tiếp theo.",
        [
            ("FR-CAL-01", "Chuyên viên khai báo/đồng bộ khung giờ khả dụng và loại hình Online/Offline.", "Must", "Chỉ slot hợp lệ được chọn; múi giờ hiển thị Asia/Ho_Chi_Minh."),
            ("FR-CAL-02", "Đặt lịch bắt buộc có Student, chuyên viên, ngày giờ, thời lượng, hình thức và địa điểm/link.", "Must", "Thiếu dữ liệu bị chặn với thông báo rõ."),
            ("FR-CAL-03", "Ngăn lịch trùng của cùng chuyên viên; override chỉ dành cho người có quyền và phải ghi lý do.", "Must", "Hai lịch giao nhau bị chặn; override xuất hiện trong audit."),
            ("FR-CAL-04", "Gửi xác nhận và nhắc lịch theo template; kênh và thời điểm là cấu hình.", "Must", "Log cho biết gửi thành công/thất bại; retry có kiểm soát."),
            ("FR-CAL-05", "Cho phép xác nhận, đổi lịch, hủy, no-show và hoàn thành; lưu lý do.", "Must", "Timeline case cập nhật ngay và không mất lịch sử cũ."),
            ("FR-CAL-06", "Biên bản tư vấn lưu tóm tắt, nhu cầu, ngân sách, mục tiêu, kết luận, next action và mức độ riêng tư.", "Must", "Không thể hoàn thành buổi tư vấn nếu thiếu kết luận và next action."),
            ("FR-CAL-07", "Tự tạo task follow-up từ biên bản và gán người phụ trách/deadline.", "Must", "Task xuất hiện trên worklist đúng owner."),
            ("FR-CAL-08", "Đồng bộ lịch với Microsoft 365/Google Calendar theo cấu hình.", "Should", "Có external event ID, reconciliation và xử lý lỗi đồng bộ."),
            ("FR-CAL-09", "Hỗ trợ lịch nhóm/hội thảo và giới hạn số chỗ.", "Could", "Không nhận đăng ký vượt capacity; có danh sách tham dự."),
        ],
        [
            "Thời lượng slot và buffer giữa hai buổi là cấu hình.",
            "Thay đổi lịch sau khi đã gửi thông báo phải phát sinh thông báo cập nhật.",
            "Nội dung biên bản chỉ hiển thị theo quyền; BGH mặc định xem tổng hợp.",
        ],
    )

    add_requirement_module(
        doc, "4.4", "Student 360 và quản lý case", "Tạo một màn hình trung tâm giúp chuyên viên hiểu hồ sơ và hành động tiếp theo mà không phải mở nhiều nguồn.",
        [
            ("FR-STU-01", "Header hiển thị ảnh, Mã HS, lớp, niên khóa, liên hệ, CMHS, owner, trạng thái và badge rủi ro.", "Must", "Dữ liệu cốt lõi nhất quán với danh sách và API."),
            ("FR-STU-02", "Hiển thị hành trình và % tiến độ theo milestone cấu hình.", "Must", "% chỉ tính từ milestone active; có giải thích cách tính."),
            ("FR-STU-03", "Tổng hợp KPI: trường mục tiêu, application, deadline, file mới, item thiếu và rủi ro.", "Must", "KPI drill-down đến danh sách nguồn; không lệch module chuyên sâu."),
            ("FR-STU-04", "Quản lý thông tin cá nhân, địa chỉ, liên hệ khẩn cấp và thông tin CMHS có validation.", "Must", "Thay đổi trường nhạy cảm có audit; dữ liệu trống hiển thị 'Chưa cập nhật'."),
            ("FR-STU-05", "Quản lý quốc gia, ngành, trường mục tiêu, intake, học bổng, ngân sách và ưu tiên.", "Must", "Cho phép nhiều mục tiêu và đánh dấu primary/backup."),
            ("FR-STU-06", "Hiển thị timeline hợp nhất từ đăng ký, lịch, note, upload, review, application và kết quả.", "Must", "Có bộ lọc theo loại sự kiện và sắp xếp thời gian."),
            ("FR-STU-07", "Lưu ghi chú nội bộ với mức Public-to-team / Private / Restricted.", "Must", "Ghi chú Restricted không lộ ở portal/export thường."),
            ("FR-STU-08", "Lưu lịch sử phân công và cho phép bàn giao case kèm checklist.", "Must", "Owner mới nhận task mở; owner cũ mất quyền theo cấu hình."),
            ("FR-STU-09", "Đóng, tạm dừng và mở lại case với lý do, ngày review và người phê duyệt.", "Must", "Case đóng không nhận thay đổi nghiệp vụ thường; reopen có audit."),
            ("FR-STU-10", "Tìm kiếm toàn cục theo Mã HS, tên, email, điện thoại, trường mục tiêu và mã application.", "Must", "Kết quả tuân thủ record scope và trả về trong NFR-PERF-02."),
            ("FR-STU-11", "Cho phép pin hồ sơ, lưu recent items và mở nhanh module liên quan.", "Could", "Không ảnh hưởng quyền; danh sách pin theo từng người dùng."),
        ],
        [
            "Một học sinh chỉ có một Student record; có thể có nhiều case theo năm/intake nếu nghiệp vụ chốt cần.",
            "Không lặp cột thông tin cá nhân trong bảng đang ở ngữ cảnh một học sinh.",
            "Mọi badge rủi ro phải có lý do và nguồn dữ liệu để chuyên viên giải thích được.",
        ],
        ("UI-03_student_360.png", "Hình 2 - UI-03: Student 360 Overview (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.5", "Học thuật và điểm chuẩn hóa", "Lưu điểm học kỳ, điểm môn và nhiều lần thi chuẩn hóa cùng minh chứng.",
        [
            ("FR-ACA-01", "Lưu năm học, khối, học kỳ/giữa kỳ/cuối kỳ, điểm trung bình và điểm từng môn cho lớp 10-12.", "Must", "Không tạo hai bản active cùng kỳ nếu không có version."),
            ("FR-ACA-02", "Hỗ trợ thang điểm 10 và GPA 4.0; công thức quy đổi là cấu hình có phiên bản.", "Must", "Hiển thị cả điểm nguồn và điểm quy đổi; truy được công thức áp dụng."),
            ("FR-ACA-03", "Nhập tay hoặc import bảng điểm có preview, mapping cột và kiểm tra miền giá trị.", "Must", "Dòng sai được chỉ rõ ô, lý do; không ghi dữ liệu lỗi."),
            ("FR-ACA-04", "Lưu nhiều lần thi IELTS, SAT, TOEFL, AP, HSK và loại thi mở rộng.", "Must", "Mỗi lần thi có ngày, tổng điểm, subscore, loại và file minh chứng."),
            ("FR-ACA-05", "Đánh dấu kết quả tốt nhất/mới nhất theo rule; không xóa lần thi cũ.", "Must", "Người dùng chuyển chế độ hiển thị Best/Latest/All."),
            ("FR-ACA-06", "Theo dõi ngày hết hạn hoặc ngưỡng sử dụng của chứng chỉ nếu áp dụng.", "Should", "Tạo cảnh báo trước ngưỡng cấu hình; không tự kết luận tính hợp lệ trường."),
            ("FR-ACA-07", "Liên kết điểm/chứng chỉ với file evidence đã duyệt và applications sử dụng.", "Must", "Mở được file theo quyền; application biết điểm nào đang khai báo."),
            ("FR-ACA-08", "Hiển thị xu hướng học tập và khoảng trống dữ liệu theo năm/học kỳ.", "Should", "Thiếu kỳ được cảnh báo; biểu đồ không làm sai dữ liệu nguồn."),
            ("FR-ACA-09", "Khóa/chốt kỳ điểm đã xác minh; sửa sau khóa cần quyền và lý do.", "Must", "Giá trị trước/sau có audit; người thường không sửa được."),
        ],
        [
            "Điểm rỗng khác điểm 0; UI và import không được quy đổi lẫn nhau.",
            "Đơn vị nghiệp vụ quyết định công thức GPA; hệ thống không tự giả định quy tắc của từng trường đại học.",
            "File transcript chính thức có thể mang cờ Restricted.",
        ],
    )

    add_requirement_module(
        doc, "4.6", "Evidence: chứng chỉ, giải thưởng, CLB và hoạt động", "Tạo kho minh chứng có cấu trúc, review được và tự động gán folder.",
        [
            ("FR-EVI-01", "Evidence item lưu category, tiêu đề, đơn vị tổ chức, vai trò, cấp độ, ngày, mô tả, thành tích và tag.", "Must", "Category bắt buộc; trường theo category có validation."),
            ("FR-EVI-02", "Một item có một hoặc nhiều file/ảnh/link; file có checksum, MIME, kích thước và uploader.", "Must", "Không mất metadata khi đổi tên hiển thị."),
            ("FR-EVI-03", "Học sinh/chuyên viên upload qua drag-drop và thấy tiến độ, lỗi, retry.", "Must", "Upload lỗi không tạo item hoàn tất giả; retry không tạo file trùng."),
            ("FR-EVI-04", "Tự gán folder theo mapping category; hỗ trợ hàng đợi ngoại lệ khi không xác định được.", "Must", "Folder đúng 100% với mapping test; ngoại lệ không bị thất lạc."),
            ("FR-EVI-05", "Chuẩn hóa tên file theo Mã HS - Category - Title - Date - Version, xử lý ký tự không hợp lệ.", "Must", "Tên file an toàn; tên gốc vẫn lưu trong metadata."),
            ("FR-EVI-06", "Review workflow: Pending, Approved, Needs Supplement, Rejected; có comment.", "Must", "Mỗi chuyển trạng thái có người, thời gian, lý do."),
            ("FR-EVI-07", "Phân biệt completeness (Đủ/Thiếu) và review status.", "Must", "Hai trường độc lập trong data model, filter và báo cáo."),
            ("FR-EVI-08", "Bulk assign category/reviewer và bulk request supplement theo quyền.", "Should", "Preview số item; xử lý một phần có báo cáo lỗi."),
            ("FR-EVI-09", "Tạo Portfolio Summary có chọn lọc evidence để dùng khi viết CV/essay.", "Should", "Export chứa đúng item được chọn và link/file được phép."),
            ("FR-EVI-10", "Quét file độc hại và chặn loại file/dung lượng ngoài allow-list.", "Must", "File chưa quét không được tải xuống bởi người dùng thường."),
        ],
        [
            "Mỗi evidence có đúng một category chính và có thể có nhiều tag phụ.",
            "Tác vụ đồng bộ storage phải có trạng thái Pending/Success/Failed và retry có giới hạn.",
            "Xóa evidence là soft-delete; file vật lý xử lý theo chính sách lưu giữ.",
        ],
        ("UI-04_evidence.png", "Hình 3 - UI-04: Evidence Vault và auto-folder (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.7", "University Application Pipeline", "Theo dõi từng nguyện vọng từ shortlist đến kết quả và nhập học.",
        [
            ("FR-APP-01", "Một application đại diện duy nhất cho Student + University + Program + Intake.", "Must", "Hệ thống cảnh báo bản ghi trùng; cho phép nhiều campus/program khác nhau."),
            ("FR-APP-02", "Lưu quốc gia, trường, chương trình, campus, intake, vòng nộp và priority Reach/Target/Safety.", "Must", "Trường bắt buộc được validate theo status."),
            ("FR-APP-03", "Theo dõi nhiều deadline: application, scholarship, document, interview, deposit và visa.", "Must", "Cột nearest deadline lấy mốc active sớm nhất và drill-down được."),
            ("FR-APP-04", "Quản lý workflow chuẩn bị, ready, submitted, interview, awaiting, result, withdrawn, enrolled.", "Must", "Chuyển trạng thái tuân thủ từ điển trạng thái và có audit."),
            ("FR-APP-05", "Lưu yêu cầu đầu vào tại thời điểm tư vấn như IELTS/SAT/GPA và link nguồn.", "Should", "Có ngày cập nhật; không ghi đè điểm thật của học sinh."),
            ("FR-APP-06", "Liên kết checklist và document items áp dụng cho application.", "Must", "Không cho Ready/Submitted nếu gate bắt buộc chưa đạt, trừ override có lý do."),
            ("FR-APP-07", "Ghi ngày nộp, mã hồ sơ ngoài, kênh nộp và xác nhận nộp.", "Must", "Submitted bắt buộc có ngày; mã ngoài là searchable."),
            ("FR-APP-08", "Quản lý phỏng vấn, người tham gia, link/địa điểm, kết quả và follow-up.", "Must", "Lịch xuất hiện ở worklist và calendar; có nhắc hẹn."),
            ("FR-APP-09", "Ghi Result độc lập Scholarship: Offer, Conditional Offer, Waitlist, Reject, Withdrawn.", "Must", "Có ngày kết quả, file, điều kiện offer và deadline phản hồi."),
            ("FR-APP-10", "Lưu học bổng theo %, giá trị, tiền tệ, năm áp dụng và điều kiện duy trì.", "Must", "Không dùng một trường text duy nhất; tổng hợp quy đổi theo rule báo cáo."),
            ("FR-APP-11", "So sánh shortlist theo chi phí, ranking nhập tay/nguồn, yêu cầu, deadline và priority.", "Should", "Không coi dữ liệu tham khảo là dữ liệu chính thức nếu thiếu nguồn/ngày."),
            ("FR-APP-12", "Chọn application ưu tiên và ghi quyết định trường nhập học.", "Must", "Tối đa một Enrolled active trên một intake, trừ quyền override."),
        ],
        [
            "Priority chiến lược không phải application status.",
            "Application có kết quả nhưng không có học bổng là trạng thái hợp lệ.",
            "Deadline quá hạn không tự đóng; phải có người xử lý và reason code.",
        ],
        ("UI-05_application_pipeline.png", "Hình 4 - UI-05: University Application Pipeline (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.8", "Documents, Essays, Recommendation Letters và Checklist", "Quản lý từng hạng mục tài liệu, phiên bản, người phụ trách, review và độ đầy đủ.",
        [
            ("FR-DOC-01", "Tạo document item từ checklist template theo quốc gia/trường/program hoặc thủ công.", "Must", "Item thừa hưởng category, owner role, bắt buộc, deadline và review gate."),
            ("FR-DOC-02", "Một item có thể áp dụng General hoặc nhiều applications; lưu quan hệ many-to-many.", "Must", "Bảng hiển thị rút gọn nhưng chi tiết liệt kê đủ trường áp dụng."),
            ("FR-DOC-03", "Quản lý trạng thái làm việc và trạng thái duyệt độc lập.", "Must", "Progress và ReviewStatus có filter, audit và báo cáo riêng."),
            ("FR-DOC-04", "Mỗi upload tạo DocumentVersion với version number, editor, timestamp, note, checksum và file link.", "Must", "Version cũ mở được; không ghi đè âm thầm."),
            ("FR-DOC-05", "Essay hỗ trợ comment/review theo phiên bản, assignee, due date và resolved status.", "Must", "Comment gắn đúng version; final không mất lịch sử comment."),
            ("FR-DOC-06", "Cho phép đánh dấu Final/Approved và khóa version; mở khóa cần quyền và lý do.", "Must", "Version final không bị thay thế; tạo version mới nếu chỉnh sửa."),
            ("FR-DOC-07", "Recommendation Letter hỗ trợ chế độ confidential/blind; học sinh chỉ thấy trạng thái.", "Must", "File không xuất hiện trong portal, API hoặc export của học sinh."),
            ("FR-DOC-08", "Có thể gửi link upload có thời hạn cho giáo viên/người giới thiệu.", "Should", "Link một lần/giới hạn item, hết hạn, có revoke và audit."),
            ("FR-DOC-09", "Tính checklist completion theo active required items; hiển thị % và x/y theo nhóm.", "Must", "Công thức giải thích được; item N/A không tính mẫu số."),
            ("FR-DOC-10", "Tạo cảnh báo thiếu, quá hạn, sắp đến hạn, waiting review và rejected.", "Must", "Cảnh báo nhất quán UI-01/UI-03/UI-07 và task engine."),
            ("FR-DOC-11", "Cho phép export application package theo trường với manifest file.", "Should", "Chỉ gồm version approved/final, đúng quyền và ghi log export."),
            ("FR-DOC-12", "Preview PDF/ảnh và tải file; Office file mở qua provider nếu được hỗ trợ.", "Should", "Không tải file chưa quét; link có thời hạn."),
            ("FR-DOC-13", "Không xóa vật lý item/version từ giao diện nghiệp vụ.", "Must", "Soft-delete có reason, audit và quy trình khôi phục."),
        ],
        [
            "Owner có thể là Học sinh, CMHS, Giáo viên, Nhà trường hoặc Chuyên viên.",
            "Chuyển 'Needs Revision' phải có comment hành động được và thông báo cho owner.",
            "Checklist template có version; case đang chạy không bị thay đổi hàng loạt nếu chưa xác nhận migration.",
        ],
        ("UI-07_documents_essays.png", "Hình 5 - UI-07: Documents, Essays & Checklist (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.9", "Task, cảnh báo và thông báo", "Biến sự kiện thành hành động có người chịu trách nhiệm và deadline.",
        [
            ("FR-TSK-01", "Tạo task thủ công với title, category, case, owner, priority, due date, status và dependency.", "Must", "Task thiếu owner/due date bị cảnh báo theo rule."),
            ("FR-TSK-02", "Rule engine tạo task/cảnh báo từ deadline, file mới, review, lịch hẹn, trạng thái và dữ liệu thiếu.", "Must", "Rule có version, bật/tắt, ngưỡng và log lần chạy."),
            ("FR-TSK-03", "Worklist sắp xếp quá hạn > deadline gần > file mới > rủi ro > ưu tiên thủ công.", "Must", "Thứ tự test được và người dùng hiểu lý do ưu tiên."),
            ("FR-TSK-04", "Hỗ trợ complete, reopen, reassign, snooze và escalate với lý do.", "Must", "Không mất lịch sử; snooze không che task quá deadline bắt buộc."),
            ("FR-TSK-05", "Thông báo trong app và email; Teams/SMS theo cấu hình.", "Must", "Mỗi thông báo có status queued/sent/failed/read và link an toàn."),
            ("FR-TSK-06", "Hỗ trợ digest hằng ngày/tuần và quiet hours cho cảnh báo không khẩn cấp.", "Should", "Critical bỏ qua quiet hours theo chính sách; tránh gửi trùng."),
            ("FR-TSK-07", "Template thông báo có biến, preview, version và ngôn ngữ vi-VN.", "Must", "Thiếu biến không gửi; log lưu template version."),
            ("FR-TSK-08", "Email không chứa file hoặc dữ liệu nhạy cảm; dùng link yêu cầu đăng nhập.", "Must", "Kiểm thử không rò nội dung confidential trong email."),
        ],
        [
            "Một sự kiện chỉ tạo một task active theo idempotency key trừ khi rule cho phép lặp.",
            "Escalation không tự thay owner nghiệp vụ; chỉ bổ sung người theo dõi trừ khi rule nói rõ.",
            "Người dùng có thể cấu hình kênh nhận trong giới hạn chính sách bắt buộc.",
        ],
        ("UI-01_worklist.png", "Hình 6 - UI-01: Worklist ưu tiên hằng ngày của chuyên viên (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.10", "Dashboard, báo cáo và export", "Cung cấp góc nhìn vận hành cho chuyên viên và góc nhìn quản trị cho Trưởng phòng/BGH.",
        [
            ("FR-RPT-01", "Dashboard chuyên viên hiển thị case cần xử lý, deadline 7/14/30 ngày, file mới, task quá hạn và lịch hôm nay.", "Must", "Mọi widget drill-down đúng danh sách và bộ lọc."),
            ("FR-RPT-02", "Dashboard Trưởng phòng hiển thị workload, aging, pipeline, rủi ro, completeness và kết quả theo chuyên viên.", "Must", "Tổng số khớp dữ liệu chi tiết; có định nghĩa KPI."),
            ("FR-RPT-03", "BGH xem KPI tổng hợp theo năm học, khối, quốc gia, intake và cohort.", "Must", "Không mặc định lộ dữ liệu cá nhân; drill-down theo quyền."),
            ("FR-RPT-04", "Hỗ trợ bộ lọc dùng chung, lưu bộ lọc cá nhân và reset về mặc định.", "Must", "Saved filter không chia sẻ dữ liệu vượt quyền."),
            ("FR-RPT-05", "Xuất Excel/CSV theo bộ lọc, timezone, locale và quyền cột.", "Must", "File có timestamp, người xuất, filter summary; export được audit."),
            ("FR-RPT-06", "Báo cáo kết quả gồm offer, scholarship, enrollment, country, university, program và counselor.", "Must", "Có cohort definition và tránh đếm trùng học sinh/application."),
            ("FR-RPT-07", "Cho phép lịch gửi báo cáo tổng hợp không nhạy cảm đến nhóm được phê duyệt.", "Should", "Có owner, lịch, người nhận, preview và revoke."),
            ("FR-RPT-08", "Cung cấp dataset/view cho Power BI hoặc công cụ BI qua tài khoản read-only.", "Should", "Schema versioned; dữ liệu tuân thủ scope/ẩn danh theo use case."),
        ],
        [
            "KPI phải có owner, công thức, grain, filter và thời điểm làm mới.",
            "Không cộng học bổng nhiều tiền tệ nếu chưa có quy tắc quy đổi và ngày tỷ giá được phê duyệt.",
            "Số lượng Student, Case và Application là ba grain khác nhau; mọi báo cáo phải ghi rõ grain.",
        ],
    )

    add_requirement_module(
        doc, "4.11", "Portal học sinh và CMHS", "Cho phép tự phục vụ có kiểm soát, giảm trao đổi rời rạc nhưng không lộ ghi chú nội bộ.",
        [
            ("FR-POR-01", "Học sinh xem trạng thái case, milestone, next actions, lịch hẹn, checklist và thông báo của mình.", "Must", "Không thấy ghi chú nội bộ, risk nội bộ hoặc hồ sơ khác."),
            ("FR-POR-02", "Học sinh cập nhật trường cho phép; thay đổi nhạy cảm vào hàng đợi duyệt.", "Must", "Giá trị cũ giữ đến khi duyệt; người review thấy diff."),
            ("FR-POR-03", "Upload evidence/document vào item hoặc category; hiển thị trạng thái quét và review.", "Must", "File được gắn đúng StudentID và item; lỗi có hướng dẫn."),
            ("FR-POR-04", "Học sinh xác nhận/đổi lịch/hủy theo policy và thấy link online đúng thời điểm.", "Must", "Thay đổi cập nhật calendar và thông báo."),
            ("FR-POR-05", "CMHS liên kết với học sinh qua quy trình xác minh/đồng ý và quyền cấu hình.", "Must", "Thu hồi liên kết có hiệu lực ngay; có audit."),
            ("FR-POR-06", "Cho phép xác nhận đã đọc yêu cầu và trả lời request supplement.", "Should", "Chuyên viên thấy timestamp và phản hồi trong timeline."),
            ("FR-POR-07", "Giao diện mobile responsive cho các thao tác thiết yếu.", "Must", "Hoạt động tại width 360px; upload và checklist sử dụng được."),
        ],
        [
            "Portal không hiển thị file Recommendation Letter confidential.",
            "Mọi link deep-link yêu cầu đăng nhập và kiểm tra quyền tại thời điểm mở.",
            "Dữ liệu liên hệ do CMHS sửa không tự ghi đè SIS nếu chưa có quy trình đồng bộ.",
        ],
    )

    add_requirement_module(
        doc, "4.12", "Website công khai và CMS", "Duy trì kênh truyền thông, học bổng, tuyển sinh và CTA đăng ký tư vấn; phạm vi đề xuất Phase 3.",
        [
            ("FR-WEB-01", "Website hiển thị banner, tin tức, học bổng, tuyển sinh, câu chuyện học sinh/cựu học sinh và hành trình.", "Should", "Chỉ nội dung Published và PublishAt <= hiện tại hiển thị."),
            ("FR-WEB-02", "Mọi category/bài có thể cấu hình CTA tới form tư vấn hoặc landing page.", "Should", "CTA đúng URL, có tracking source và không hiển thị khi để trống."),
            ("FR-WEB-03", "Tìm kiếm, tag, category, bài nổi bật và responsive desktop/mobile.", "Should", "Không trả bài nháp; trang lỗi/empty state rõ ràng."),
            ("FR-CMS-01", "CMS tạo bài với title, summary, rich content, category, tag, author, cover, gallery và SEO fields.", "Should", "Thiếu title/category/content không publish được."),
            ("FR-CMS-02", "Workflow Draft, In Review, Scheduled, Published, Archived; quyền publish riêng.", "Should", "Mọi chuyển trạng thái có audit; schedule chạy đúng múi giờ."),
            ("FR-CMS-03", "Preview trước publish và version history/restore.", "Should", "Preview không public; restore tạo version mới."),
            ("FR-CMS-04", "Media library kiểm tra định dạng, kích thước, alt text và quyền sử dụng.", "Should", "Ảnh thiếu alt bị cảnh báo; file ngoài allow-list bị chặn."),
            ("FR-CMS-05", "Form tư vấn từ website tạo CounselingRegistration và lưu source/campaign.", "Must", "Không tạo đường dữ liệu riêng ngoài module đăng ký."),
        ],
        [
            "Website public phải tách quyền và bề mặt tấn công khỏi hệ thống nội bộ ở mức kiến trúc phù hợp.",
            "Nội dung chưa publish không được lộ qua API, sitemap hoặc URL đoán được.",
            "Dữ liệu form public phải có chống spam/rate limit và thông báo riêng tư.",
        ],
        ("PW-01_public_website.png", "Hình 7 - PW-01: Website công khai của Phòng HTQT (nguồn UI v2.0)."),
    )

    add_requirement_module(
        doc, "4.13", "Quản trị hệ thống và master data", "Cho phép IT và chủ nghiệp vụ cấu hình hệ thống mà không sửa code cho các thay đổi thường xuyên.",
        [
            ("FR-ADM-01", "Quản lý năm học, khối, quốc gia, trường, program/major, intake và loại deadline.", "Must", "Không xóa master đang được tham chiếu; dùng inactive/effective date."),
            ("FR-ADM-02", "Quản lý document type, evidence category, checklist template và folder mapping có version.", "Must", "Preview tác động trước publish; case cũ không đổi ngoài ý muốn."),
            ("FR-ADM-03", "Quản lý workflow/status dictionary, transition và reason code trong giới hạn thiết kế.", "Should", "Không tạo transition mâu thuẫn; cấu hình có validation."),
            ("FR-ADM-04", "Quản lý ngưỡng cảnh báo, risk rule, SLA và notification template.", "Must", "Có test/preview rule và rollback version."),
            ("FR-ADM-05", "Quản lý user-role-scope; xem access history và export danh sách quyền.", "Must", "Thay đổi nhạy cảm yêu cầu quyền riêng và audit."),
            ("FR-ADM-06", "Audit viewer lọc theo thời gian, user, student, entity, action, IP/request ID.", "Must", "Không cho sửa/xóa log từ UI thường."),
            ("FR-ADM-07", "Integration console hiển thị job, last success, error, retry, reconciliation và dead-letter.", "Must", "Lỗi không bị che; retry không tạo bản ghi trùng."),
            ("FR-ADM-08", "Cấu hình retention, export policy, file allow-list và feature flag theo môi trường.", "Should", "Thay đổi có phê duyệt/ghi log; không áp dụng ngược âm thầm."),
            ("FR-ADM-09", "Cung cấp health/status dashboard cho dịch vụ, storage, queue, email và lịch.", "Should", "Hiển thị dependency, thời điểm kiểm tra và trạng thái."),
            ("FR-ADM-10", "Hỗ trợ soft-delete restore cho entity được phép trong thời hạn cấu hình.", "Must", "Khôi phục giữ ID/quan hệ; mọi thao tác có audit."),
        ],
        [
            "Master data dùng effective date để bảo toàn dữ liệu lịch sử.",
            "Admin kỹ thuật và admin nội dung là hai vai trò khác nhau.",
            "Cấu hình production phải có backup/version và cơ chế rollback.",
        ],
        ("PW-02_cms.png", "Hình 8 - PW-02: CMS / Admin đăng bài (nguồn UI v2.0)."),
    )


def build_data_model(doc: Document) -> None:
    add_heading(doc, "5. MÔ HÌNH DỮ LIỆU VÀ QUẢN TRỊ TÀI LIỆU", 1)
    add_heading(doc, "5.1. Nguyên tắc dữ liệu", 2)
    for item in (
        "StudentID là khóa nghiệp vụ; khóa kỹ thuật dùng ID bất biến, không chứa thông tin cá nhân.",
        "Mọi entity quan trọng có CreatedAt, CreatedBy, UpdatedAt, UpdatedBy, Version và trạng thái soft-delete nếu áp dụng.",
        "Dữ liệu có grain rõ: Student, Case, Registration, Application, DocumentItem và FileVersion không được gộp tùy tiện.",
        "Thời gian lưu UTC ở backend; hiển thị Asia/Ho_Chi_Minh; ngày không có giờ dùng kiểu date.",
        "File nằm ở storage provider; database giữ metadata, checksum, quyền, trạng thái quét và provider reference.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "5.2. Data dictionary cấp entity", 2)
    entities = [
        ("Student", "StudentId", "StudentCode, FullName, DOB, Grade, Class, AcademicYear, Contacts, Status", "1-n Registration, Case, Score, Evidence"),
        ("StudentGuardian", "GuardianId", "StudentId, Name, Relation, Phone, Email, ConsentStatus", "n-1 Student"),
        ("CounselingRegistration", "RegistrationId", "StudentId?, Source, Interests, PreferredSlots, Status, DuplicateFlag", "n-1 Student; 1-n Appointment"),
        ("CounselingAppointment", "AppointmentId", "RegistrationId/StudentId, CounselorId, StartAt, Duration, Mode, Location, Status", "1-1/1-n Session"),
        ("CounselingSession", "SessionId", "AppointmentId, Summary, Conclusion, NextAction, PrivacyLevel", "n-1 Student; 1-n Task"),
        ("StudentCase", "CaseId", "StudentId, SchoolYear, TargetIntake, Stage, RiskLevel, OwnerId, OpenedAt, ClosedAt", "n-1 Student"),
        ("CounselorAssignment", "AssignmentId", "CaseId, CounselorId, Role, StartAt, EndAt, HandoverNote", "n-1 Case/User"),
        ("StudyTarget", "TargetId", "CaseId, Country, University?, Program?, Priority, Budget, ScholarshipGoal", "n-1 Case"),
        ("AcademicRecord", "AcademicRecordId", "StudentId, SchoolYear, GradeLevel, TermType, Average, GPA4, FormulaVersion", "1-n SubjectScore"),
        ("SubjectScore", "SubjectScoreId", "AcademicRecordId, SubjectCode, Score, Scale, Status", "n-1 AcademicRecord"),
        ("TestScore", "TestScoreId", "StudentId, TestType, TestDate, Total, Subscores, ExpiryDate, EvidenceId", "n-1 Student"),
        ("EvidenceItem", "EvidenceId", "StudentId, Category, Title, Role, Level, Dates, Completeness, ReviewStatus", "1-n FileVersion/Tag"),
        ("StorageObject", "StorageObjectId", "Provider, ExternalId, FolderId, FileName, Mime, Size, Checksum, ScanStatus", "n-1 Evidence/DocumentVersion"),
        ("University", "UniversityId", "Name, Country, Campus, Active, Source, UpdatedAt", "1-n Program/Application"),
        ("Program", "ProgramId", "UniversityId, Name, DegreeLevel, Campus, Active", "1-n Application"),
        ("UniversityApplication", "ApplicationId", "CaseId, UniversityId, ProgramId, Intake, Round, Priority, Status, Result", "1-n Deadline/Interview/Scholarship"),
        ("ApplicationDeadline", "DeadlineId", "ApplicationId, Type, DueAt, Status, Source, OwnerId", "n-1 Application"),
        ("Interview", "InterviewId", "ApplicationId, ScheduleAt, Mode, Location, Status, Outcome", "n-1 Application"),
        ("Scholarship", "ScholarshipId", "ApplicationId, Name, ValueType, Amount, Currency, Percent, Conditions", "n-1 Application"),
        ("DocumentItem", "DocumentId", "CaseId, Type, Category, OwnerRole/User, Required, Progress, ReviewStatus, DueAt", "n-n Application; 1-n Version"),
        ("DocumentVersion", "DocumentVersionId", "DocumentId, VersionNo, StorageObjectId, EditorId, Note, IsFinal, ApprovedAt", "n-1 DocumentItem"),
        ("ReviewComment", "CommentId", "DocumentVersionId/EvidenceId, AuthorId, Text, Status, CreatedAt", "n-1 Version/Item"),
        ("ChecklistItem", "ChecklistItemId", "CaseId/ApplicationId, TemplateVersion, DocumentId?, Status, Required, DueAt", "n-1 Case/Application"),
        ("Task", "TaskId", "CaseId, EntityRef, Type, Title, OwnerId, Priority, DueAt, Status, RuleRunId", "n-1 Case/User"),
        ("Notification", "NotificationId", "RecipientId, Channel, TemplateVersion, EntityRef, Status, SentAt, ReadAt", "n-1 User/Student"),
        ("PublicPost", "PostId", "Category, Title, Summary, ContentRef, Status, PublishAt, CTAUrl", "n-1 Author/Media"),
        ("AuditEvent", "AuditId", "ActorId, Action, EntityType, EntityId, BeforeHash/Value, AfterHash/Value, At, RequestId", "Append-only"),
        ("IntegrationJob", "JobId", "Type, CorrelationId, StartedAt, CompletedAt, Status, RetryCount, ErrorCode", "1-n job item"),
    ]
    add_table(doc, ["Entity", "Khóa", "Trường chính", "Quan hệ chính"], entities, [1800, 1450, 4410, 1700], alt_rows=True, font_size=7.8)
    add_heading(doc, "5.3. Phân loại dữ liệu", 2)
    add_table(
        doc,
        ["Mức", "Ví dụ", "Kiểm soát tối thiểu"],
        [
            ("Public", "Bài viết đã xuất bản, thông tin liên hệ công khai.", "Public read; CMS publish workflow."),
            ("Internal", "Danh mục trường, template, cấu hình không nhạy cảm.", "Đăng nhập; quyền theo vai trò."),
            ("Sensitive", "Thông tin cá nhân, điểm, hồ sơ ứng tuyển, liên hệ CMHS.", "RBAC + record scope; mã hóa; audit export/download."),
            ("Restricted", "Thư giới thiệu confidential, tài chính, ghi chú riêng tư.", "Need-to-know; quyền riêng; link ngắn hạn; audit bắt buộc."),
        ],
        [1500, 3400, 4460],
    )
    add_heading(doc, "5.4. Quy tắc file", 2)
    for item in (
        "Allow-list mặc định: PDF, DOCX, XLSX, PPTX, JPG, PNG; danh sách và dung lượng tối đa là cấu hình.",
        "MIME thực tế phải khớp extension; file được quét trước khi preview/download.",
        "Checksum dùng phát hiện file trùng; không dùng checksum để tự hợp nhất hai item nghiệp vụ.",
        "Tên hiển thị có thể thay đổi; provider ID và checksum là bất biến cho version.",
        "Link chia sẻ không public, có hạn sử dụng và ràng buộc người dùng nếu provider hỗ trợ.",
    ):
        add_bullet(doc, item)


def build_ui(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "6. YÊU CẦU GIAO DIỆN VÀ TRẢI NGHIỆM", 1)
    add_heading(doc, "6.1. Danh mục màn hình", 2)
    add_table(
        doc,
        ["Mã", "Màn hình", "Người dùng chính", "Mục tiêu"],
        [
            ("UI-01", "Counselor Workbench", "Chuyên viên / Trưởng phòng", "Biết việc cần làm hôm nay, deadline, file mới và rủi ro."),
            ("UI-02", "Counseling Registration Queue", "Tiếp nhận / Trưởng phòng", "Chống trùng, phân công, đặt lịch."),
            ("UI-03", "Student 360", "Chuyên viên", "Toàn cảnh học sinh và điều hướng module."),
            ("UI-04", "Evidence Vault", "Chuyên viên / Học sinh", "Chứng chỉ, hoạt động, review và auto-folder."),
            ("UI-05", "Application Pipeline", "Chuyên viên / Trưởng phòng", "Trường, ngành, deadline, kết quả và học bổng."),
            ("UI-06", "Academic & Test Scores", "Chuyên viên", "Điểm học kỳ, điểm môn và standardized tests."),
            ("UI-07", "Documents, Essays & Checklist", "Chuyên viên / cộng tác", "Version, review, owner và completeness."),
            ("UI-08", "Calendar & Counseling Sessions", "Chuyên viên / Tiếp nhận", "Lịch, xung đột, biên bản, follow-up."),
            ("UI-09", "Reports & Operations", "Trưởng phòng / BGH", "KPI, workload, pipeline, outcome."),
            ("POR-01", "Student/Parent Portal", "Học sinh / CMHS", "Tự phục vụ, upload, lịch và next actions."),
            ("PW-01", "Public Website", "Khách truy cập", "Nội dung và CTA đăng ký."),
            ("CMS-01", "Content Management", "Admin CMS", "Soạn, duyệt, publish và media."),
        ],
        [1000, 2550, 2300, 3510],
        alt_rows=True,
        font_size=8.6,
        center_columns=(0,),
    )
    add_heading(doc, "6.2. Chuẩn tương tác", 2)
    for item in (
        "Mọi danh sách lớn có search, filter, sort, pagination, tổng số kết quả, reset và lưu bộ lọc.",
        "Bulk action chỉ bật khi có lựa chọn hợp lệ; phải preview số bản ghi và kết quả xử lý một phần.",
        "Các trạng thái Loading, Empty, Error, No permission và Partial data phải có thông điệp và hành động tiếp theo.",
        "Màu không là tín hiệu duy nhất; luôn có nhãn/icon/text cho risk, status và validation.",
        "Form dài chia section; autosave chỉ dùng nơi không gây hiểu lầm và phải hiển thị trạng thái lưu.",
        "Điều hướng rời trang khi có thay đổi chưa lưu phải cảnh báo.",
        "Tên học sinh và Mã HS luôn xuất hiện ở header khi thao tác trong một hồ sơ.",
        "Ngày hiển thị dd/MM/yyyy; giờ 24h; timezone Asia/Ho_Chi_Minh; dữ liệu tiền có currency rõ.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "6.3. Kích thước và responsive", 2)
    add_table(
        doc,
        ["Bề mặt", "Yêu cầu"],
        [
            ("Nội bộ desktop", "Tối ưu từ 1366x768; bảng cho phép sticky header/column, resize hợp lý và horizontal scroll có kiểm soát."),
            ("Tablet", "Đọc và thao tác đơn giản; bảng phức tạp có chế độ card hoặc chọn cột."),
            ("Portal mobile", "Sử dụng được từ width 360px cho login, checklist, lịch, upload và thông báo."),
            ("Website public", "Responsive, ảnh tối ưu, điều hướng bàn phím và metadata chia sẻ."),
        ],
        [2300, 7060],
    )
    add_heading(doc, "6.4. Accessibility và ngôn ngữ", 2)
    for item in (
        "Mục tiêu WCAG 2.1 AA cho luồng trọng yếu; bàn phím sử dụng được và focus visible.",
        "Label liên kết với input; lỗi nêu rõ trường, nguyên nhân và cách sửa.",
        "Tương phản màu đủ; text có thể zoom 200% mà không mất chức năng trọng yếu.",
        "Ngôn ngữ mặc định vi-VN; dữ liệu Unicode đầy đủ; kiến trúc cho phép thêm tiếng Anh.",
    ):
        add_bullet(doc, item)


def build_integrations(doc: Document) -> None:
    add_heading(doc, "7. TÍCH HỢP VÀ API", 1)
    add_heading(doc, "7.1. Nguyên tắc", 2)
    for item in (
        "SRS không ép framework; API phải versioned, kiểm tra quyền, validate input và trả error code nhất quán.",
        "Mọi write API hỗ trợ correlation/request ID; thao tác có nguy cơ lặp dùng idempotency key.",
        "Tích hợp bất đồng bộ có retry với backoff, dead-letter và reconciliation; không retry vô hạn.",
        "Không đưa file lớn qua API business nếu storage provider hỗ trợ upload trực tiếp có URL ký.",
        "Dữ liệu tích hợp có data owner, source of truth, tần suất và quy tắc xung đột.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "7.2. Danh mục tích hợp", 2)
    add_table(
        doc,
        ["Hệ thống", "Dữ liệu / chiều", "Tần suất", "Xử lý lỗi"],
        [
            ("Identity / SSO", "User, group, status -> hệ thống", "Real-time login + sync định kỳ", "Từ chối an toàn; log; cảnh báo admin."),
            ("SIS", "Student, class, guardian -> hệ thống; status phản hồi tùy chọn", "Ngày hoặc near-real-time", "Staging, validation, reject report, reconciliation."),
            ("Drive / OneDrive / SharePoint", "Folder, file, permission <-> hệ thống", "Event + background sync", "Retry idempotent; orphan queue; health alert."),
            ("Calendar", "Appointment <-> calendar event", "Near-real-time", "External ID, retry, conflict log, reconciliation."),
            ("Email / Teams", "Notification -> người dùng", "Event-driven", "Queue, retry, bounce/error status, fallback in-app."),
            ("Power BI / BI", "Curated dataset -> BI", "Theo lịch", "Last refresh, schema version, quality check."),
            ("Website/CMS", "CTA/form -> Registration", "Real-time", "Rate limit, anti-spam, retry và dedupe."),
        ],
        [1800, 3300, 1900, 2360],
        alt_rows=True,
        font_size=8.4,
    )
    add_heading(doc, "7.3. Nhóm API nghiệp vụ tối thiểu", 2)
    add_table(
        doc,
        ["Nhóm", "Năng lực"],
        [
            ("Students/Cases", "Search, get 360, update profile, assign, transition, timeline."),
            ("Registrations/Counseling", "Submit, dedupe, verify, assign, schedule, session notes, conclusion."),
            ("Academic/Test", "CRUD có version/lock, import preview/commit, best/latest query."),
            ("Evidence/Files", "Create item, upload session, finalize, review, supplement, folder sync."),
            ("Applications", "CRUD, transition, deadlines, interview, result, scholarship, enrollment."),
            ("Documents/Checklist", "Template instantiate, item/version, review, comment, package export."),
            ("Tasks/Notifications", "Query worklist, update task, rule event, preferences, read status."),
            ("Reports/Admin", "KPI query, export job, master/version, audit search, integration health."),
        ],
        [2400, 6960],
        font_size=9,
    )


def build_nfr(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "8. YÊU CẦU PHI CHỨC NĂNG", 1)
    add_callout(doc, "Mô hình tải chuẩn", "Tối thiểu 1.000 case/năm, 5.000 case tích lũy, 100 người dùng nội bộ đồng thời và 300 phiên portal đồng thời tại cao điểm; IT xác nhận lại bằng sizing trước build.", "info")
    nfr_rows = [
        ("NFR-PERF-01", "Trang danh sách/360 tải p95 <= 3 giây, p99 <= 5 giây với dữ liệu chuẩn, không tính file preview ngoài.", "Must"),
        ("NFR-PERF-02", "Tìm kiếm toàn cục p95 <= 2 giây; filter/pagination p95 <= 2,5 giây.", "Must"),
        ("NFR-PERF-03", "Save nghiệp vụ p95 <= 2 giây; thao tác dài chuyển background job và hiển thị tiến độ.", "Must"),
        ("NFR-PERF-04", "Dashboard tổng hợp p95 <= 5 giây hoặc dùng dữ liệu cache có thời điểm làm mới rõ.", "Must"),
        ("NFR-PERF-05", "Export 10.000 dòng hoàn thành <= 60 giây; chạy nền và thông báo khi sẵn sàng.", "Should"),
        ("NFR-SCL-01", "Mở rộng tối thiểu 5.000 case, 100.000 document/evidence metadata mà không đổi mô hình lõi.", "Must"),
        ("NFR-AVL-01", "Uptime mục tiêu 99,5% theo tháng trong giờ làm việc, trừ bảo trì đã thông báo.", "Must"),
        ("NFR-DR-01", "RPO <= 24 giờ và RTO <= 8 giờ cho baseline; file dùng version/backup theo provider và chính sách IT.", "Must"),
        ("NFR-SEC-01", "Mã hóa khi truyền bằng HTTPS/TLS hiện hành; mã hóa dữ liệu lưu trữ theo năng lực nền tảng được duyệt.", "Must"),
        ("NFR-SEC-02", "Không có lỗ hổng mức Critical/High chưa chấp nhận rủi ro tại thời điểm go-live.", "Must"),
        ("NFR-SEC-03", "Upload được kiểm tra MIME, allow-list, size, malware và quyền tải xuống.", "Must"),
        ("NFR-PRV-01", "Dữ liệu cá nhân và Restricted tuân thủ policy nhà trường, mục đích sử dụng, quyền truy cập và lưu giữ đã phê duyệt.", "Must"),
        ("NFR-AUD-01", "100% sự kiện audit bắt buộc có actor, action, entity, timestamp, result và request ID.", "Must"),
        ("NFR-USE-01", "Người dùng nghiệp vụ hoàn thành 10 tác vụ cốt lõi sau đào tạo ngắn; thông báo lỗi có hướng dẫn sửa.", "Must"),
        ("NFR-ACC-01", "Luồng trọng yếu đạt mục tiêu WCAG 2.1 AA và sử dụng bằng bàn phím.", "Should"),
        ("NFR-COMP-01", "Hỗ trợ hai phiên bản ổn định gần nhất của Chrome/Edge; Safari cho portal nếu trường yêu cầu.", "Must"),
        ("NFR-LOC-01", "vi-VN, Unicode, dd/MM/yyyy, giờ 24h, Asia/Ho_Chi_Minh; tiền tệ luôn có currency.", "Must"),
        ("NFR-OBS-01", "Có log, metric, trace/request ID, health check, cảnh báo lỗi tích hợp và dashboard vận hành.", "Must"),
        ("NFR-MNT-01", "Cấu hình, schema, API và rule có version; migration có rollback/forward plan.", "Must"),
        ("NFR-QLT-01", "Kiểm tra dữ liệu tự động phát hiện orphan file, duplicate key, invalid status và deadline mâu thuẫn.", "Should"),
    ]
    add_table(doc, ["ID", "Yêu cầu đo được", "Ưu tiên"], nfr_rows, [1200, 7160, 1000], alt_rows=True, center_columns=(0, 2), font_size=8.45)
    add_heading(doc, "8.1. Tiêu chí kiểm thử hiệu năng", 2)
    for item in (
        "Dữ liệu test phải phản ánh phân bố thật: case active/closed, file metadata, task, timeline và application.",
        "Báo cáo p50/p95/p99, throughput, error rate và tài nguyên; không chỉ báo thời gian trung bình.",
        "Test tối thiểu search, Student 360, worklist, upload metadata, save application, dashboard và export.",
        "Kiểm thử khi một dependency chậm/lỗi để xác nhận timeout, retry và thông báo người dùng.",
    ):
        add_bullet(doc, item)


def build_security(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "9. BẢO MẬT, RIÊNG TƯ VÀ AUDIT", 1)
    add_heading(doc, "9.1. Kiểm soát bắt buộc", 2)
    add_table(
        doc,
        ["Vùng kiểm soát", "Yêu cầu"],
        [
            ("Identity", "SSO/MFA theo IdP; session timeout; revoke; account lifecycle; service account riêng."),
            ("Authorization", "RBAC + record scope + confidential flag tại backend; deny by default."),
            ("Data", "Mã hóa truyền/lưu; secrets không nằm trong code/log; masking dữ liệu nhạy cảm khi phù hợp."),
            ("Files", "Malware scan, MIME/size, signed URL ngắn hạn, provider permission, checksum và quarantine."),
            ("Application", "Validate input, chống injection, CSRF phù hợp, rate limit, secure headers và error không lộ nội bộ."),
            ("Export", "Quyền riêng, cột theo scope, audit, thời hạn file, watermark/notice nếu cần."),
            ("Operations", "Tách DEV/UAT/PROD; production access giới hạn; change/backup/incident process."),
            ("Privacy", "Mục đích sử dụng, consent khi cần, quyền truy cập, lưu giữ và xử lý yêu cầu dữ liệu theo policy."),
        ],
        [2200, 7160],
        alt_rows=True,
        font_size=8.8,
    )
    add_heading(doc, "9.2. Danh mục sự kiện audit tối thiểu", 2)
    for item in (
        "Đăng nhập thành công/thất bại, logout, session revoke, thay đổi role/scope.",
        "Xem/download/export dữ liệu Restricted; tạo link chia sẻ; mở Recommendation Letter confidential.",
        "Tạo/sửa/xóa mềm/hợp nhất Student, Case, Application, Evidence, Document, Score và CMS post.",
        "Chuyển trạng thái, phân công, đặt/đổi/hủy lịch, kết luận tư vấn và override deadline/gate.",
        "Upload, quét, review, publish version, restore và thay đổi permission storage.",
        "Thay đổi master data, rule, template, retention, integration secret/reference và feature flag.",
        "Integration job, retry, reconciliation, import và export.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "9.3. Nội dung audit record", 2)
    add_table(
        doc,
        ["Trường", "Mô tả"],
        [
            ("Who", "Actor user/service, role, scope; support mode nếu có."),
            ("What", "Action, entity type, entity ID, field groups thay đổi."),
            ("When", "Timestamp UTC và timezone hiển thị."),
            ("Where", "IP/device/session/request ID theo policy."),
            ("Before/After", "Giá trị trước/sau hoặc hash/diff phù hợp; không log secret/plain sensitive quá mức."),
            ("Result", "Success/Denied/Failed, reason/error code."),
        ],
        [1700, 7660],
    )
    add_callout(doc, "Bất biến", "Audit log không được sửa/xóa từ giao diện nghiệp vụ. Thời hạn lưu và quyền xem log phải được BGH/IT phê duyệt.", "risk")


def build_reporting(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "10. BÁO CÁO VÀ KPI", 1)
    add_heading(doc, "10.1. KPI dictionary", 2)
    kpis = [
        ("KPI-01", "Registrations", "Số CounselingRegistration tạo trong kỳ; grain Registration."),
        ("KPI-02", "Verified Students", "Số Student duy nhất có đăng ký đã xác minh trong kỳ."),
        ("KPI-03", "Counseling conversion", "Số case Active / số đăng ký đã tư vấn đủ điều kiện; loại duplicate/cancel theo rule."),
        ("KPI-04", "Active cases", "Số StudentCase ở trạng thái active tại thời điểm báo cáo."),
        ("KPI-05", "Unassigned aging", "Thời gian từ registration hợp lệ đến assignment đầu tiên."),
        ("KPI-06", "Counselor workload", "Case active, task open, critical/high risk và lịch trong kỳ theo counselor."),
        ("KPI-07", "Checklist completion", "Completed required active items / total required active items."),
        ("KPI-08", "On-time readiness", "Application đạt Ready trước internal deadline / application có deadline."),
        ("KPI-09", "Submission rate", "Application Submitted / application trong shortlist đã chốt."),
        ("KPI-10", "Offer rate", "Applications có Offer/Conditional Offer / applications có kết quả; grain Application."),
        ("KPI-11", "Student offer rate", "Students có ít nhất một offer / students có ít nhất một application có kết quả."),
        ("KPI-12", "Scholarship count/value", "Số offer có scholarship và giá trị theo currency; không cộng khác currency nếu chưa quy đổi."),
        ("KPI-13", "Enrollment destinations", "Số Student chọn Enrolled theo country/university/program/intake."),
        ("KPI-14", "Document review SLA", "Thời gian từ upload hoàn tất đến review đầu tiên."),
        ("KPI-15", "Data completeness", "Tỷ lệ case đủ nhóm dữ liệu bắt buộc theo stage."),
    ]
    add_table(doc, ["ID", "KPI", "Định nghĩa"], kpis, [950, 2300, 6110], alt_rows=True, center_columns=(0,), font_size=8.55)
    add_heading(doc, "10.2. Bộ lọc chuẩn", 2)
    for item in (
        "Năm học, cohort, khối/lớp, intake, quốc gia, trường, program, counselor và case stage.",
        "Khoảng ngày theo event date phù hợp: registration date, submission date, result date hoặc snapshot date.",
        "Risk level, document completeness, application status, result và scholarship flag.",
        "Mọi báo cáo ghi rõ grain, timezone, thời điểm làm mới và filter đang áp dụng.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "10.3. Báo cáo bắt buộc", 2)
    add_table(
        doc,
        ["Báo cáo", "Người dùng", "Tần suất / hành động"],
        [
            ("Daily Operations", "Chuyên viên", "Hằng ngày; xử lý task, deadline, file mới và lịch."),
            ("Team Workload & Risk", "Trưởng phòng", "Hằng ngày/tuần; cân tải, escalations và aging."),
            ("Application Pipeline", "Chuyên viên/Trưởng phòng", "Theo intake; deadline, submission, interview, result."),
            ("Outcome & Scholarship", "Trưởng phòng/BGH", "Theo kỳ/năm; offer, scholarship, enrollment."),
            ("Data Quality", "Data owner/Admin", "Hằng tuần; thiếu dữ liệu, duplicate, orphan, invalid status."),
            ("Access & Audit", "IT/được ủy quyền", "Theo tháng/sự cố; access review, export, restricted access."),
        ],
        [3000, 2500, 3860],
        alt_rows=True,
    )


def build_uat(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "11. KIỂM THỬ, UAT VÀ TIÊU CHÍ NGHIỆM THU", 1)
    add_heading(doc, "11.1. Nguyên tắc nghiệm thu", 2)
    for item in (
        "Mỗi Must FR có ít nhất một test case pass và liên kết bằng ID.",
        "Không còn defect Critical/High mở; Medium có workaround và kế hoạch được chủ nghiệp vụ chấp nhận.",
        "UAT dùng dữ liệu ẩn danh/giả lập nhưng đủ độ phức tạp, không dùng hồ sơ thật tùy tiện.",
        "Security, backup/restore, permission và integration failure là một phần nghiệm thu, không chỉ kiểm tra UI happy path.",
        "Tài liệu vận hành, hướng dẫn người dùng, data dictionary và cấu hình baseline được bàn giao cùng release.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "11.2. Kịch bản UAT trọng yếu", 2)
    uat = [
        ("UAT-01", "Đăng ký mới có Mã HS hợp lệ", "Liên kết/tạo đúng Student; có registration ID; không trùng."),
        ("UAT-02", "Đăng ký lặp cùng Mã HS", "Chỉ một Student; lịch sử có hai registrations; cảnh báo/link đúng."),
        ("UAT-03", "Thiếu Mã HS và nghi trùng", "Needs Verification; không tự hợp nhất; reviewer thấy candidate."),
        ("UAT-04", "Bulk assign chuyên viên", "Owner và audit cập nhật; người ngoài scope không thấy case."),
        ("UAT-05", "Đặt hai lịch trùng", "Lịch thứ hai bị chặn; override có quyền và reason mới thành công."),
        ("UAT-06", "Hoàn thành tư vấn", "Có conclusion, next action; case stage/task/timeline cập nhật."),
        ("UAT-07", "Student 360", "KPI, điểm, evidence, app, tài liệu và cảnh báo khớp module nguồn."),
        ("UAT-08", "Import bảng điểm có dòng lỗi", "Preview chỉ rõ lỗi; commit dòng hợp lệ; có kết quả import."),
        ("UAT-09", "Nhiều lần thi IELTS/SAT", "Hiển thị Best/Latest/All đúng rule; giữ mọi lần thi."),
        ("UAT-10", "Upload evidence hợp lệ", "Quét pass; folder đúng mapping; metadata/audit đầy đủ."),
        ("UAT-11", "Upload file nguy hiểm/sai MIME", "Quarantine/reject; không preview/download; có log lỗi."),
        ("UAT-12", "Request supplement evidence", "Status/comment/task/notification cập nhật cho đúng owner."),
        ("UAT-13", "Tạo application trùng", "Cảnh báo theo Student + University + Program + Intake."),
        ("UAT-14", "Deadline <= 7 ngày", "Xuất hiện worklist, Student 360, pipeline và thông báo nhất quán."),
        ("UAT-15", "Chuyển Submitted khi thiếu gate", "Bị chặn; override có quyền và lý do được audit."),
        ("UAT-16", "Ghi offer và scholarship", "Hai dữ liệu độc lập; KPI/report cập nhật đúng grain."),
        ("UAT-17", "Upload essay version mới", "Tạo vN+1; version cũ/comment còn; final lock đúng."),
        ("UAT-18", "Recommendation confidential", "Học sinh chỉ thấy status; không mở file qua URL/API/export."),
        ("UAT-19", "Checklist completion", "% và x/y đúng; item N/A không vào mẫu số."),
        ("UAT-20", "Bàn giao case", "Owner mới có task/quyền; owner cũ mất scope theo policy; lịch sử giữ."),
        ("UAT-21", "Export theo bộ lọc", "Chỉ đúng dòng/cột được phép; có timestamp/filter/audit."),
        ("UAT-22", "CMHS truy cập sai học sinh", "Bị từ chối; không lộ metadata; có security log."),
        ("UAT-23", "Storage/Email tạm lỗi", "Queue/retry; người dùng thấy trạng thái; không nhân bản dữ liệu."),
        ("UAT-24", "Khôi phục soft-delete/version", "Giữ ID/quan hệ; audit đầy đủ; quyền được kiểm tra."),
        ("UAT-25", "Load test chuẩn", "Đạt p95/p99/error rate theo NFR; có báo cáo tài nguyên."),
        ("UAT-26", "Backup và restore drill", "Khôi phục trong RTO, mất dữ liệu không vượt RPO; checklist ký xác nhận."),
    ]
    add_table(doc, ["ID", "Kịch bản", "Kết quả mong đợi"], uat, [1000, 3500, 4860], alt_rows=True, center_columns=(0,), font_size=8.35)
    add_heading(doc, "11.3. Exit criteria", 2)
    add_table(
        doc,
        ["Nhóm", "Điều kiện ra"],
        [
            ("Chức năng", "100% Must pass; >= 95% Should pass hoặc có kế hoạch được duyệt."),
            ("Defect", "0 Critical/High mở; Medium được đánh giá rủi ro và có owner/date."),
            ("Hiệu năng", "Đạt NFR ở môi trường tương đương production với dataset chuẩn."),
            ("Bảo mật", "Security review hoàn tất; không còn finding Critical/High chưa chấp nhận."),
            ("Dữ liệu", "Reconciliation import đạt ngưỡng; duplicate/orphan được xử lý hoặc có danh sách ngoại lệ."),
            ("Vận hành", "Monitoring, backup, restore, runbook, support contact và training sẵn sàng."),
        ],
        [1900, 7460],
    )


def build_deployment(doc: Document) -> None:
    add_heading(doc, "12. TRIỂN KHAI, DI CHUYỂN DỮ LIỆU VÀ VẬN HÀNH", 1)
    add_heading(doc, "12.1. Môi trường", 2)
    add_table(
        doc,
        ["Môi trường", "Mục đích", "Dữ liệu"],
        [
            ("DEV", "Phát triển, unit/integration test.", "Synthetic; không dùng production dump thô."),
            ("TEST/UAT", "System test và nghiệp vụ nghiệm thu.", "Ẩn danh/giả lập; tích hợp sandbox."),
            ("PROD", "Vận hành chính thức.", "Dữ liệu thật; access/change/backup chặt."),
        ],
        [1600, 3400, 4360],
    )
    add_heading(doc, "12.2. Di chuyển dữ liệu", 2)
    migration_steps = (
        "Chốt source inventory: Excel, SIS, Drive/OneDrive, lịch, danh mục và hồ sơ giấy số hóa.",
        "Lập mapping source-to-target, rule chuẩn hóa Mã HS, tên, ngày, status và master data.",
        "Chạy dry-run vào staging; xuất lỗi, duplicate candidates, orphan files và reconciliation counts.",
        "Chủ nghiệp vụ làm sạch/duyệt ngoại lệ; lặp dry-run đến khi đạt ngưỡng.",
        "Freeze window; backup nguồn; chạy migration cuối; đối soát count/hash/sample và ký xác nhận.",
        "Giữ rollback plan và bản sao read-only nguồn theo policy; không xóa nguồn ngay sau cutover.",
    )
    for item in migration_steps:
        add_number(doc, item)
    add_heading(doc, "12.3. Checklist go-live", 2)
    for item in (
        "Baseline SRS/UAT được phê duyệt; owner và support matrix được công bố.",
        "SSO, role, record scope, confidential access và service account được kiểm tra.",
        "Master data, checklist template, folder mapping, risk rules và notification template đã version/chốt.",
        "Migration và reconciliation hoàn tất; issue ngoại lệ có owner.",
        "Backup/restore drill, monitoring/alerting, incident/runbook và capacity test đã pass.",
        "Đào tạo chuyên viên, Trưởng phòng, Admin và tài liệu quick guide hoàn tất.",
        "Kế hoạch hypercare, kênh hỗ trợ, SLA và lịch retrospective được chốt.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "12.4. Vận hành và hỗ trợ", 2)
    add_table(
        doc,
        ["Sự kiện", "Mức", "Mục tiêu phản hồi", "Ví dụ"],
        [
            ("P1", "Critical", "Theo SLA IT đã phê duyệt", "Không đăng nhập diện rộng, mất dữ liệu, lộ dữ liệu, hệ thống dừng."),
            ("P2", "High", "Theo SLA IT đã phê duyệt", "Module trọng yếu lỗi, deadline/notification sai diện rộng."),
            ("P3", "Medium", "Trong giờ hỗ trợ", "Lỗi có workaround, ảnh hưởng nhóm nhỏ."),
            ("P4", "Low/Request", "Theo backlog", "Cải tiến UI, báo cáo mới, thay đổi cấu hình."),
        ],
        [1000, 1400, 2700, 4260],
        center_columns=(0, 1),
    )


def build_roadmap(doc: Document) -> None:
    add_heading(doc, "13. ROADMAP VÀ PHẠM VI PHÁT HÀNH", 1)
    add_table(
        doc,
        ["Giai đoạn", "Phạm vi", "Cổng ra"],
        [
            ("Phase 0 - Discovery & Design", "Chốt OD-01..06, process, data model, architecture, prototype, backlog và UAT plan.", "Baseline được ký; estimate và release plan được duyệt."),
            ("Phase 1 - Core CRM", "SSO/RBAC, Registration, dedupe, assignment, calendar, session, Student 360, task/worklist, audit.", "Quản lý end-to-end từ đăng ký đến Active Case."),
            ("Phase 2 - Portfolio & Applications", "Academic/test, Evidence, storage, Application pipeline, Documents/Essays/Checklist, reports.", "Vận hành hồ sơ apply hoàn chỉnh cho 1.000 case."),
            ("Phase 3 - Portal, Website & CMS", "Portal học sinh/CMHS, public website, CMS, CTA và external recommender upload.", "Tự phục vụ và truyền thông tích hợp."),
            ("Phase 4 - BI & Optimization", "SIS nâng cao, BI, workload/risk optimization, data quality automation, optional AI assist.", "KPI quản trị và tối ưu vận hành."),
        ],
        [2000, 5080, 2280],
        alt_rows=True,
        font_size=8.7,
    )
    add_heading(doc, "13.1. MVP khuyến nghị", 2)
    add_callout(doc, "MVP", "Phase 1 + phần tối thiểu của Phase 2 gồm Academic/Test, Evidence upload + auto-folder, Application, Documents/Checklist. Đây là phạm vi nhỏ nhất đáp ứng đúng nhu cầu quản lý 1.000 hồ sơ thay vì chỉ là CRM tiếp nhận.", "success")
    add_heading(doc, "13.2. Điều kiện đưa AI vào giai đoạn sau", 2)
    for item in (
        "AI chỉ hỗ trợ gợi ý tag, tóm tắt, phát hiện thiếu và tìm kiếm; không tự gửi/nộp hồ sơ hoặc quyết định thay chuyên viên.",
        "Dữ liệu đưa vào mô hình phải có phê duyệt, kiểm soát quyền, logging và chính sách không dùng huấn luyện ngoài mục đích.",
        "Mọi output AI có nhãn, nguồn và bước human review; đo độ chính xác trước rollout.",
        "Không dùng AI để tạo nhận xét gây bất lợi cho học sinh mà thiếu tiêu chí minh bạch và quy trình khiếu nại.",
    ):
        add_bullet(doc, item)


def build_appendices(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "PHỤ LỤC A - TỪ ĐIỂN TRẠNG THÁI", 1)
    add_heading(doc, "A.1. Document progress", 2)
    add_table(
        doc,
        ["Trạng thái", "Ý nghĩa", "Trạng thái kế hợp lệ"],
        [
            ("Not Started", "Chưa bắt đầu.", "In Progress, Not Applicable"),
            ("In Progress", "Đang chuẩn bị.", "Ready for Review, Blocked"),
            ("Ready for Review", "Owner đã gửi review.", "Completed, Needs Revision"),
            ("Needs Revision", "Cần chỉnh sửa theo phản hồi.", "In Progress, Ready for Review"),
            ("Blocked", "Bị chặn có lý do/dependency.", "In Progress"),
            ("Completed", "Đã hoàn tất về nội dung.", "In Progress nếu reopen có lý do"),
            ("Not Applicable", "Không áp dụng cho case/application.", "Not Started nếu bật lại"),
        ],
        [1900, 3600, 3860],
        alt_rows=True,
    )
    add_heading(doc, "A.2. Review status", 2)
    add_table(
        doc,
        ["Trạng thái", "Ý nghĩa", "Ai chuyển"],
        [
            ("Not Reviewed", "Chưa vào hàng đợi duyệt.", "Hệ thống/Owner"),
            ("Pending Review", "Đang chờ reviewer.", "Owner/Hệ thống"),
            ("Approved", "Đạt yêu cầu tại version hiện tại.", "Reviewer có quyền"),
            ("Needs Supplement", "Thiếu minh chứng/thông tin.", "Reviewer"),
            ("Needs Revision", "Cần chỉnh sửa nội dung.", "Reviewer"),
            ("Rejected", "Không chấp nhận; bắt buộc reason.", "Reviewer có quyền"),
        ],
        [2000, 4800, 2560],
    )
    add_heading(doc, "A.3. Application status", 2)
    add_table(
        doc,
        ["Trạng thái", "Điều kiện tối thiểu"],
        [
            ("Shortlisted", "Có University, Program, Intake, Priority."),
            ("Preparing", "Có checklist và owner; đang chuẩn bị."),
            ("Ready", "Gate bắt buộc đạt hoặc override được phê duyệt."),
            ("Submitted", "Có SubmittedAt và proof/external ID nếu áp dụng."),
            ("Interview", "Có interview active."),
            ("Awaiting Result", "Đã nộp/phỏng vấn; chưa có final result."),
            ("Result Received", "Có Result và ResultAt."),
            ("Withdrawn", "Có reason/date/actor."),
            ("Enrolled", "Học sinh chốt nhập học; có decision date."),
        ],
        [2600, 6760],
        alt_rows=True,
    )

    add_page_break(doc)
    add_heading(doc, "PHỤ LỤC B - MA TRẬN TRUY VẾT", 1)
    add_table(
        doc,
        ["Mục tiêu", "Nhóm yêu cầu", "Kịch bản UAT", "KPI/NFR"],
        [
            ("Không trùng hồ sơ", "FR-REG-01..10", "UAT-01..03", "KPI-01..03"),
            ("Tư vấn có lịch và kết luận", "FR-CAL-01..09", "UAT-04..06", "NFR-PERF-03"),
            ("Toàn cảnh 1 học sinh", "FR-STU-01..11", "UAT-07,20", "NFR-PERF-01..02"),
            ("Dữ liệu học thuật tin cậy", "FR-ACA-01..09", "UAT-08..09", "KPI-15"),
            ("Evidence đúng folder", "FR-EVI-01..10", "UAT-10..12", "NFR-SEC-03"),
            ("Không bỏ lỡ deadline", "FR-APP-01..12; FR-TSK-01..08", "UAT-13..16", "KPI-08; NFR-PERF-01"),
            ("Essay/tài liệu có version", "FR-DOC-01..13", "UAT-17..19", "KPI-07,14"),
            ("Quyền và riêng tư", "FR-AUT; FR-POR; mục 9", "UAT-18,22", "NFR-SEC/PRV/AUD"),
            ("Quản trị 1.000 hồ sơ", "FR-RPT; FR-ADM", "UAT-21,25", "KPI-04..15; NFR-SCL"),
            ("Vận hành bền vững", "FR-ADM; mục 7,12", "UAT-23,24,26", "NFR-AVL/DR/OBS/MNT"),
        ],
        [2500, 2700, 1900, 2260],
        alt_rows=True,
        font_size=8.5,
    )
    add_callout(doc, "Quản lý thay đổi", "Khi thêm/sửa một Must FR, đội dự án phải cập nhật test case, data mapping, quyền, audit và tài liệu đào tạo liên quan trước khi phê duyệt release.", "warning")

    add_heading(doc, "PHỤ LỤC C - CHECKLIST HỒ SƠ MẪU", 1)
    add_table(
        doc,
        ["Nhóm", "Hạng mục mẫu", "Owner mặc định", "Gate"],
        [
            ("Academic", "Transcript/Học bạ; bảng điểm; school profile", "Nhà trường/Chuyên viên", "Required"),
            ("Tests", "IELTS/TOEFL; SAT/ACT; AP/HSK theo trường", "Học sinh", "Conditional"),
            ("Essays", "Personal Statement; Common App Essay; Supplemental Essays", "Học sinh + Chuyên viên", "Required"),
            ("Recommendation", "Counselor Letter; Teacher Letters", "GV/Chuyên viên", "Required/Confidential"),
            ("Identity", "Passport; ảnh; thông tin cá nhân", "Học sinh/CMHS", "Required"),
            ("Financial", "Financial statement; sponsorship; scholarship forms", "CMHS", "By school/country"),
            ("Application", "Form; fee/proof; submission receipt", "Học sinh + Chuyên viên", "Required"),
            ("Interview", "Lịch; preparation; feedback", "Học sinh + Chuyên viên", "If applicable"),
            ("Result", "Offer/reject/waitlist; scholarship; deposit; enrollment", "Chuyên viên", "Outcome"),
        ],
        [1800, 4000, 2100, 1460],
        alt_rows=True,
        font_size=8.5,
    )
    add_body(doc, "Checklist thực tế phải được version theo quốc gia/trường/chương trình/intake. Bảng trên chỉ là seed template để IT thiết kế cấu trúc, không phải danh sách pháp lý cố định.")

    end = doc.add_paragraph()
    end.paragraph_format.space_before = Pt(24)
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = end.add_run("HẾT TÀI LIỆU - SRS v3.0")
    set_run_font(er, size=12, bold=True, color=COLORS["navy"])


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "SRS - Hệ thống Quản lý Hồ sơ Du học NSHM v3.0"
    props.subject = "Software Requirements Specification"
    props.author = "Phòng Hợp tác Quốc tế - Trường Ngôi Sao Hoàng Mai"
    props.keywords = "SRS, du học, hồ sơ, tư vấn, NSHM"
    props.comments = "Baseline bàn giao cho đội IT"


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    configure_styles(doc)
    add_custom_numbering(doc)
    set_core_properties(doc)

    build_cover(doc)
    build_front_matter(doc)
    build_overview(doc)
    build_scope_roles(doc)
    build_workflows(doc)
    build_functional_requirements(doc)
    build_data_model(doc)
    build_ui(doc)
    build_integrations(doc)
    build_nfr(doc)
    build_security(doc)
    build_reporting(doc)
    build_uat(doc)
    build_deployment(doc)
    build_roadmap(doc)
    build_appendices(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
